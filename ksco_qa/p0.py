"""P0 — 사전 검사·매핑·파일 프로필 (코드) (설계 5.1, 7.1 / 00_Work_Run.md 3장 / CLAUDE.md M1 결정 1~10).

결정적 단계다. 같은 입력이면 같은 결과를 낸다. 판정을 만들지 않고 상태·경고만 기록한다.
"""

from __future__ import annotations

import csv
import json
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from ksco_qa.config import Config
from ksco_qa.profiler import FileProfile, combine_for_row, profile_file, row_budget
from ksco_qa.profiler import fulltext
from ksco_qa.profiler.zip_file import ZIP_LOCATOR_SEP, member_display_name, profile_member

P0_VERSION = "2"
DATALIST_SUFFIX = "_DataList.xlsx"
SHEET_NAME = "DataList"
HEADER_ROW = 1
OUTPUT_FILES_DIR = "01_output_files"
# 여러 형식 행: 06 matched_locator는 문자열 하나라 경로를 " | "로 잇고, 전체 목록은 matched_locators[]에 둔다
MATCHED_LOCATOR_SEP = " | "
# 결정 3: B~I열만 검증 (설계 5.1-2의 7개 열)
EXPECTED_HEADERS = {"B": "id", "C": "직업명", "D": "산출물", "E": "레퍼런스", "F": "출력포맷", "G": "태스크", "I": "직무정보"}
DATA_COLUMNS = "BCDEFGHI"  # 교훈 #6: B~I 중 하나라도 값이 있으면 데이터 행
SHEET_ERROR_RE = re.compile(r"^#(REF!|N/A|VALUE!|NAME\?|DIV/0!|NUM!|NULL!|SPILL!|CALC!|ERROR!|GETTING_DATA)")
JOB_SEPARATOR_RE = re.compile(r"[:：]")
FORMAT_TOKEN_RE = re.compile(r"[A-Za-z0-9]+")

# 설계 7.1 row_card 필드 (+ 결정: original_task 없음 표시). 이전 검수 결과 열(R~V 등)은 절대 넣지 않는다.
ROW_CARD_FIELDS = (
    "source_row", "id", "job", "job_leaf", "deliverable_title", "format", "expected_filename",
    "match_status", "inspected", "reference_mode", "reference_note", "original_task", "original_task_empty",
    "jobinfo", "file_profile",
)
# FILE_ACCESS=fulltext일 때만 뒤에 붙는 필드 (2026-10-23). profile 모드의 row_card는 위 필드 그대로다.
FULLTEXT_FIELDS = ("file_fulltext", "file_fulltext_note")
MAPPING_AUDIT_COLUMNS = (
    "source_row", "id", "name_D", "format_F", "expected_filename", "match_status", "matched_locator",
    "inspected", "inspected_locations", "reason", "reference_mode", "warnings",
)


class P0Stop(Exception):
    """실행 중단 사유 (DataList 0개·2개 이상, 시트·헤더 불일치, id 누락·중복, NOT_FOUND 비율 초과)."""


# --- 1. DataList 찾기 ---


def find_datalist(task_folder: Path) -> tuple[Path, str]:
    """5.1-1: *_DataList.xlsx를 정확히 1개 찾고 TASK_NAME을 파일명에서 얻는다. 0개·2개 이상이면 중단.
    Excel 잠금 파일(~$...)은 세지 않는다."""
    task_folder = Path(task_folder)
    if not task_folder.is_dir():
        raise P0Stop(f"Task 폴더가 없다: {task_folder}")
    found = sorted(
        p for p in task_folder.glob(f"*{DATALIST_SUFFIX}") if p.is_file() and not p.name.startswith("~$")
    )
    if len(found) != 1:
        names = ", ".join(p.name for p in found) or "없음"
        raise P0Stop(f"*{DATALIST_SUFFIX}는 정확히 1개여야 한다 (찾은 파일: {names})")
    return found[0], found[0].name[: -len(DATALIST_SUFFIX)]


# --- 2~4. DataList 읽기 ---


def nfc(s: str) -> str:
    return unicodedata.normalize("NFC", s)


def norm_header(value: Any) -> str:
    """결정 3: NFC + 앞뒤 공백·줄바꿈 제거."""
    return nfc(str(value)).strip() if value is not None else ""


def is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _formula(cell: Any) -> str | None:
    if getattr(cell, "data_type", None) != "f":
        return None
    v = cell.value
    return v if isinstance(v, str) else getattr(v, "text", None) or str(v)


def _hyperlink(cell: Any) -> str | None:
    link = getattr(cell, "hyperlink", None)
    if link is None:
        return None
    return link.target or (f"#{link.location}" if link.location else None)


@dataclass
class DataRow:
    source_row: int
    values: dict[str, Any]  # 열 문자 → 캐시 계산값
    cells: dict[str, dict[str, Any]]  # source.cells: 열 문자 → {header, value, formula, hyperlink}


@dataclass
class DataList:
    path: Path
    sheet: str
    headers: dict[str, str | None]
    rows: list[DataRow]


def validate_headers(headers: dict[str, Any]) -> None:
    """5.1-2 / 결정 3: B~I열 헤더가 정확히 일치해야 한다. 다르면 중단."""
    wrong = [
        f"{col}열 '{norm_header(headers.get(col))}' (기대 '{name}')"
        for col, name in EXPECTED_HEADERS.items()
        if norm_header(headers.get(col)) != name
    ]
    if wrong:
        raise P0Stop("DataList 헤더(1행)가 다르다: " + "; ".join(wrong))


def read_datalist(path: Path) -> DataList:
    """5.1-3·4: 캐시 계산값으로 읽고, 다시 읽어 수식·하이퍼링크 원문을 셀별로 보관.
    데이터 행 = B~I열 중 하나라도 값이 있는 행."""
    import openpyxl
    from openpyxl.utils import get_column_letter

    wb_v = openpyxl.load_workbook(path, data_only=True)
    wb_f = openpyxl.load_workbook(path, data_only=False)
    if SHEET_NAME not in wb_v.sheetnames:
        raise P0Stop(f"'{SHEET_NAME}' 탭이 없다 (있는 탭: {', '.join(wb_v.sheetnames)})")
    ws_v, ws_f = wb_v[SHEET_NAME], wb_f[SHEET_NAME]
    max_col = max(ws_v.max_column, ws_f.max_column)
    letters = [get_column_letter(c) for c in range(1, max_col + 1)]
    headers = {L: ws_v[f"{L}{HEADER_ROW}"].value for L in letters}
    validate_headers(headers)

    rows = []
    for r in range(HEADER_ROW + 1, max(ws_v.max_row, ws_f.max_row) + 1):
        values = {L: ws_v[f"{L}{r}"].value for L in letters}
        if all(is_blank(values.get(L)) for L in DATA_COLUMNS):
            continue
        cells = {}
        for L in letters:
            fcell = ws_f[f"{L}{r}"]
            cells[L] = {
                "header": None if headers[L] is None else str(headers[L]),
                "value": json_value(values[L]),
                "formula": _formula(fcell),
                "hyperlink": _hyperlink(fcell),
            }
        rows.append(DataRow(r, values, cells))
    return DataList(path, SHEET_NAME, {L: None if v is None else str(v) for L, v in headers.items()}, rows)


# --- 5. 식별 검사 ---


def norm_id(value: Any) -> str | None:
    if is_blank(value):
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def check_identifiers(rows: list[DataRow]) -> None:
    """5.1-5 / 결정 4: id 누락·중복이면 행 목록과 함께 중단."""
    missing = [r.source_row for r in rows if norm_id(r.values.get("B")) is None]
    seen: dict[str, list[int]] = {}
    for r in rows:
        rid = norm_id(r.values.get("B"))
        if rid is not None:
            seen.setdefault(rid, []).append(r.source_row)
    dups = {k: v for k, v in seen.items() if len(v) > 1}
    problems = []
    if missing:
        problems.append(f"id 누락 {len(missing)}행: {_rows_preview(missing)}")
    if dups:
        problems.append("id 중복: " + "; ".join(f"{k} → {_rows_preview(v)}" for k, v in list(dups.items())[:20]))
    if problems:
        raise P0Stop(" / ".join(problems) + " — 시트에서 B열 id를 값으로 채운 뒤 다시 실행")


def _rows_preview(rows: list[int], limit: int = 20) -> str:
    shown = ", ".join(str(r) for r in rows[:limit])
    return shown + (f" 외 {len(rows) - limit}행" if len(rows) > limit else "") + "행"


def jobinfo_warnings(value: Any) -> list[str]:
    """결정 5: I열이 비었거나, 시트 오류값이거나, "매칭 실패"를 포함하면 경고 (중단 안 함)."""
    if is_blank(value):
        return ["JOBINFO_EMPTY"]
    s = str(value).strip()
    out = []
    if SHEET_ERROR_RE.match(s):
        out.append("JOBINFO_SHEET_ERROR")
    if "매칭 실패" in s:
        out.append("JOBINFO_MATCH_FAILED")
    return out


def job_leaf(job: Any) -> str | None:
    """결정 10: C열을 ':'(전각 '：' 포함)로 나눈 마지막 값. 구분자가 없으면 전체."""
    if is_blank(job):
        return None
    return JOB_SEPARATOR_RE.split(str(job))[-1].strip()


# --- 6. 인벤토리 ---


@dataclass
class InventoryEntry:
    name_nfc: str  # 파일 이름(확장자 포함), NFC
    locator: str  # 01_output_files/ 기준 상대 경로 (ZIP 멤버는 "<zip>!/<멤버>")
    in_zip: bool
    accessible: bool
    path: Path  # 디스크 파일 (ZIP 멤버면 ZIP 파일)
    member: str | None = None  # ZIP 안 원래 이름
    reason: str | None = None


@dataclass
class Inventory:
    entries: list[InventoryEntry]
    broken_zips: list[str] = field(default_factory=list)  # 멤버 목록을 못 읽은 ZIP

    @property
    def incomplete(self) -> bool:
        return bool(self.broken_zips)


def _readable(path: Path) -> tuple[bool, str | None]:
    try:
        with open(path, "rb") as f:
            f.read(1)
        return True, None
    except OSError as e:
        return False, f"열 수 없음: {type(e).__name__}"


def build_inventory(output_dir: Path) -> Inventory:
    """5.1-6: 01_output_files/ 전체와 ZIP 멤버 목록화. 이름은 NFC.
    결정 8: 암호 걸린 멤버·권한 오류 = accessible False, ZIP이 깨져 목록을 못 읽으면 broken_zips."""
    output_dir = Path(output_dir)
    if not output_dir.is_dir():
        raise P0Stop(f"{OUTPUT_FILES_DIR}/ 폴더가 없다")
    entries, broken = [], []
    for path in sorted(output_dir.rglob("*"), key=lambda p: nfc(p.relative_to(output_dir).as_posix())):
        if not path.is_file():
            continue
        rel = nfc(path.relative_to(output_dir).as_posix())
        ok, reason = _readable(path)
        entries.append(InventoryEntry(nfc(path.name), rel, False, ok, path, None, reason))
        if path.suffix.lower() != ".zip" or not ok:
            continue
        try:
            with zipfile.ZipFile(path) as z:
                infos = [i for i in z.infolist() if not i.is_dir()]
        except (zipfile.BadZipFile, OSError) as e:
            broken.append(f"{rel} ({type(e).__name__})")
            continue
        for info in infos:
            display = nfc(member_display_name(info))
            encrypted = bool(info.flag_bits & 0x1)
            entries.append(InventoryEntry(
                nfc(display.rsplit("/", 1)[-1]), f"{rel}{ZIP_LOCATOR_SEP}{display}", True, not encrypted,
                path, info.filename, "암호 걸린 ZIP 멤버" if encrypted else None,
            ))
    return Inventory(entries, broken)


# --- 7. 매핑 ---


def match_key(filename: str) -> tuple[str, str]:
    """전체 이름 비교 키: NFC 이름 + 확장자 대소문자 무시 (00 3.1)."""
    name = nfc(filename)
    stem, dot, ext = name.rpartition(".")
    return (stem, ext.casefold()) if dot else (name, "")


def parse_formats(value: Any) -> tuple[list[str] | None, str | None]:
    """F열: 양끝 공백·앞 점 1개만 정리 (00 3.1). 쉼표로 여러 형식이면 각각 (결정).
    '/' 등으로 묶인 복수 형식·누락·해석 불가는 (None, 사유) → CONFLICT."""
    if is_blank(value):
        return None, "F열 누락"
    raw = str(value).strip()
    formats = []
    for token in raw.split(","):
        token = token.strip()
        if token.startswith("."):
            token = token[1:]
        if not FORMAT_TOKEN_RE.fullmatch(token):
            return None, f"F열 형식 해석 불가: {raw!r}"
        if token.casefold() not in (f.casefold() for f in formats):
            formats.append(token)
    return formats, None


@dataclass
class MatchResult:
    expected_filenames: list[str]
    match_status: str  # enums.yaml match_status
    matched: list[InventoryEntry]
    reason: str
    warnings: list[str] = field(default_factory=list)

    @property
    def expected_filename(self) -> str | None:
        return ", ".join(self.expected_filenames) or None

    @property
    def matched_locators(self) -> list[str]:
        return [e.locator for e in self.matched]


def build_index(inventory: Inventory) -> dict[tuple[str, str], list[InventoryEntry]]:
    index: dict[tuple[str, str], list[InventoryEntry]] = {}
    for e in inventory.entries:
        index.setdefault(match_key(e.name_nfc), []).append(e)
    return index


def match_deliverable(
    name_d: Any, format_f: Any, index: dict[tuple[str, str], list[InventoryEntry]], inventory_incomplete: bool
) -> MatchResult:
    """5.1-7: 예상 파일명 {D}.{F}와 전체 이름이 정확히 같은 파일 (NFC·확장자 대소문자만 허용).
    결정 6: D는 비교할 때만 앞뒤 공백 제거 + D_TRAILING_SPACE 경고.
    결정 7: 디스크와 ZIP 멤버에 같은 이름 → AMBIGUOUS. 결정 8: 열 수 없음 → INACCESSIBLE,
    ZIP이 깨져 목록이 불완전하면 NOT_FOUND 대신 INVENTORY_INCOMPLETE.
    여러 형식: 전부 찾으면 MATCHED, 일부만 MATCHED + PARTIAL_FORMAT, 하나도 없으면 NOT_FOUND."""
    warnings = []
    if is_blank(name_d):
        return MatchResult([], "CONFLICT", [], "D열 누락")
    d = str(name_d)
    if d != d.strip():
        warnings.append("D_TRAILING_SPACE")
    d = d.strip()
    formats, problem = parse_formats(format_f)
    if formats is None:
        return MatchResult([], "CONFLICT", [], problem or "F열 결함", warnings)

    expected = [f"{d}.{fmt.lower()}" for fmt in formats]
    found, ambiguous, inaccessible, missing = [], [], [], []
    for name in expected:
        cands = index.get(match_key(name), [])
        if len(cands) > 1:
            ambiguous.append(f"{name}: " + ", ".join(c.locator for c in cands))
        elif len(cands) == 1 and not cands[0].accessible:
            inaccessible.append(f"{cands[0].locator} ({cands[0].reason})")
        elif len(cands) == 1:
            found.append(cands[0])
        else:
            missing.append(name)

    if ambiguous:
        return MatchResult(expected, "AMBIGUOUS", [], "같은 이름 후보 여러 개: " + " / ".join(ambiguous), warnings)
    if inaccessible:
        return MatchResult(expected, "INACCESSIBLE", [], "파일은 있으나 열 수 없음: " + " / ".join(inaccessible), warnings)
    if found and missing:
        warnings.append("PARTIAL_FORMAT")
        return MatchResult(expected, "MATCHED", found, "일부 형식만 찾음 (없음: " + ", ".join(missing) + ")", warnings)
    if found:
        return MatchResult(expected, "MATCHED", found, "D_F_EXACT_FILENAME", warnings)
    if inventory_incomplete:
        return MatchResult(expected, "INVENTORY_INCOMPLETE", [], "파일 없음 — 멤버 목록을 못 읽은 ZIP이 있음", warnings)
    return MatchResult(expected, "NOT_FOUND", [], "완전한 인벤토리에 같은 이름 없음", warnings)


def check_not_found_ratio(results: list[MatchResult], stop_ratio: float, inventory: Inventory) -> None:
    """5.1-8: NOT_FOUND가 데이터 행의 stop_ratio를 넘으면 중단하고 경로·ZIP·권한 문제를 보고."""
    if not results:
        return
    nf = [r for r in results if r.match_status == "NOT_FOUND"]
    ratio = len(nf) / len(results)
    if ratio <= stop_ratio:
        return
    on_disk = [e for e in inventory.entries if not e.in_zip]
    in_zip = [e for e in inventory.entries if e.in_zip]
    exts: dict[str, int] = {}
    for e in inventory.entries:
        ext = match_key(e.name_nfc)[1] or "(없음)"
        exts[ext] = exts.get(ext, 0) + 1
    examples = ", ".join(r.expected_filename or "?" for r in nf[:5])
    raise P0Stop(
        f"NOT_FOUND {len(nf)}/{len(results)}행 ({ratio:.0%}) > 기준 {stop_ratio:.0%}. "
        f"인벤토리: 디스크 파일 {len(on_disk)}개, ZIP 멤버 {len(in_zip)}개, 형식 {exts}. "
        f"없는 예: {examples}. 경로·ZIP 압축 해제·권한·파일명(D/F) 문제를 먼저 확인"
    )


# --- 9. Reference ---


def classify_reference(e_value: Any) -> tuple[str, str | None]:
    """결정 1: 비어 있으면 GENERATE_REFERENCE, 값이 있으면 UNRESOLVED_REFERENCE + E열 원문.
    1차 구현은 PROVIDED_REFERENCE를 판정하지 않고 Reference 파일도 찾지 않는다."""
    if is_blank(e_value):
        return "GENERATE_REFERENCE", None
    return "UNRESOLVED_REFERENCE", str(e_value)


# --- 10. 파일 프로필 ---


def profile_matched(match: MatchResult, cache_dir: Path, total_budget: int) -> list[FileProfile]:
    budget = row_budget(match.matched_locators, total_budget)
    profiles = []
    for e in match.matched:
        if e.in_zip:
            profiles.append(profile_member(e.path, e.member, e.locator, cache_dir, budget))
        else:
            profiles.append(profile_file(e.path, e.locator, cache_dir, budget))
    return profiles


def fulltext_matched(match: MatchResult, cache_dir: Path, max_tokens: int) -> list[fulltext.FullText]:
    items = []
    for e in match.matched:
        if e.in_zip:
            items.append(fulltext.fulltext_member(e.path, e.member, e.locator, cache_dir, max_tokens))
        else:
            items.append(fulltext.fulltext_file(e.path, e.locator, cache_dir, max_tokens))
    return items


# --- 11. row_card ---


def _text(value: Any) -> str | None:
    return None if value is None else str(json_value(value))


def build_row_card(row: DataRow, match: MatchResult, inspected: bool, reference: tuple[str, str | None],
                   file_profile: str | None, full: tuple[str | None, str] | None = None) -> dict[str, Any]:
    """5.1-11, 7.1: row_card. 설계 7.1 필드만 담는다 (O·L·M열, 이전 검수 결과 열 금지)."""
    v = row.values
    original = None if is_blank(v.get("G")) else str(v["G"])
    card = {
        "source_row": row.source_row,
        "id": norm_id(v.get("B")),
        "job": _text(v.get("C")),
        "job_leaf": job_leaf(v.get("C")),
        "deliverable_title": _text(v.get("D")),
        "format": _text(v.get("F")),
        "expected_filename": match.expected_filename,
        "match_status": match.match_status,
        "inspected": inspected,
        "reference_mode": reference[0],
        "reference_note": reference[1],
        "original_task": original,
        "original_task_empty": original is None,
        "jobinfo": _text(v.get("I")),
        "file_profile": file_profile,
    }
    assert tuple(card) == ROW_CARD_FIELDS
    if full is not None:
        card["file_fulltext"], card["file_fulltext_note"] = full
    return card


# --- 출력 ---


def write_once_jsonl(path: Path, records: list[dict[str, Any]]) -> None:
    """추가 전용 원칙: 이미 있으면 내용이 같을 때만 그대로 두고, 다르면 오류."""
    text = "".join(json.dumps(r, ensure_ascii=False) + "\n" for r in records)
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        if path.read_text(encoding="utf-8") != text:
            raise P0Stop(f"{path.name}가 이전 결과와 다르다 (덮어쓰기 금지). 입력·코드가 바뀌었으면 새 run으로 실행")
        return
    path.write_text(text, encoding="utf-8")


def write_mapping_audit(path: Path, rows: list[dict[str, Any]]) -> None:
    """03_results/00_mapping_audit.csv (00 3.3 최소 열 + warnings). Excel에서 열리게 UTF-8 BOM."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=MAPPING_AUDIT_COLUMNS)
        w.writeheader()
        w.writerows(rows)


@dataclass
class P0Result:
    task_name: str
    datalist: DataList
    row_cards: list[dict[str, Any]]
    records: list[dict[str, Any]]  # P0 행 기록 (매핑 상세·경고·프로필)
    inventory: Inventory


def run_p0(task_folder: Path, cfg: Config, run_dir: Path, mapping_audit_path: Path) -> P0Result:
    """5.1 전체. 완료 조건: 데이터 행 수 = row_card 수 = mapping_audit 행 수, 모든 행에 매핑 상태."""
    task_folder, run_dir = Path(task_folder), Path(run_dir)
    datalist_path, task_name = find_datalist(task_folder)
    datalist = read_datalist(datalist_path)
    if not datalist.rows:
        raise P0Stop("데이터 행이 없다")
    check_identifiers(datalist.rows)
    inventory = build_inventory(task_folder / OUTPUT_FILES_DIR)
    index = build_index(inventory)

    matches = [match_deliverable(r.values.get("D"), r.values.get("F"), index, inventory.incomplete)
               for r in datalist.rows]
    check_not_found_ratio(matches, cfg.NOT_FOUND_STOP_RATIO, inventory)

    cache_dir = run_dir / "file_profiles"
    row_cards, records, audit = [], [], []
    for row, match in zip(datalist.rows, matches):
        profiles = profile_matched(match, cache_dir, cfg.PROFILE_MAX_CHARS) if match.match_status == "MATCHED" else []
        inspected = match.match_status == "MATCHED" and bool(profiles) and all(p.inspected for p in profiles)
        reference = classify_reference(row.values.get("E"))
        file_profile = combine_for_row(profiles) if profiles else None
        full_items: list[fulltext.FullText] = []
        full = None
        if cfg.FILE_ACCESS == "fulltext":
            if match.match_status == "MATCHED":
                full_items = fulltext_matched(match, cache_dir, cfg.FULLTEXT_MAX_TOKENS)
                full = fulltext.combine_for_row(full_items, cfg.FULLTEXT_MAX_TOKENS)
            else:
                full = (None, f"매핑 {match.match_status}: 원본 없음")
        warnings = list(match.warnings) + jobinfo_warnings(row.values.get("I"))
        warnings += list(dict.fromkeys(p.warning for p in profiles if p.warning))
        if is_blank(row.values.get("G")):
            warnings.append("ORIGINAL_TASK_EMPTY")
        reason = match.reason
        failed = [p for p in profiles if not p.inspected]
        if failed:
            reason += " / NOT_INSPECTED: " + "; ".join(f"{p.locator}: {p.error}" for p in failed)
        locations = [f"{p.locator}: {loc}" for p in profiles for loc in p.inspected_locations]

        card = build_row_card(row, match, inspected, reference, file_profile, full)
        row_cards.append(card)
        records.append({
            "source_row": row.source_row,
            "id": card["id"],
            "expected_filenames": match.expected_filenames,
            "match_status": match.match_status,
            "match_method": "D_F_EXACT_FILENAME" if match.match_status == "MATCHED" else None,
            "matched_locator": MATCHED_LOCATOR_SEP.join(match.matched_locators) or None,
            "matched_locators": match.matched_locators,
            "inspected": inspected,
            "inspected_locations": locations,
            "profiles": [p.to_dict() for p in profiles],
            **({"fulltext": [{k: v for k, v in item.to_dict().items() if k != "text"} for item in full_items],
                "file_fulltext_note": full[1]} if full is not None else {}),
            "reference_mode": reference[0],
            "reference_note": reference[1],
            "reason": reason,
            "warnings": warnings,
        })
        audit.append({
            "source_row": row.source_row,
            "id": card["id"],
            "name_D": _text(row.values.get("D")),
            "format_F": _text(row.values.get("F")),
            "expected_filename": match.expected_filename,
            "match_status": match.match_status,
            "matched_locator": MATCHED_LOCATOR_SEP.join(match.matched_locators),
            "inspected": inspected,
            "inspected_locations": "; ".join(locations),
            "reason": reason,
            "reference_mode": reference[0],
            "warnings": "; ".join(warnings),
        })

    if not (len(datalist.rows) == len(row_cards) == len(audit)):
        raise P0Stop("완료 조건 실패: 데이터 행 수·row_card 수·mapping_audit 행 수가 다르다")
    write_once_jsonl(run_dir / "row_cards.jsonl", row_cards)
    write_once_jsonl(run_dir / "source_cells.jsonl",
                     [{"source_row": r.source_row, "id": norm_id(r.values.get("B")), "cells": r.cells}
                      for r in datalist.rows])
    write_once_jsonl(run_dir / "p0" / "records.jsonl", records)
    write_mapping_audit(Path(mapping_audit_path), audit)
    return P0Result(task_name, datalist, row_cards, records, inventory)


__all__ = ["P0Stop", "P0Result", "find_datalist", "run_p0"]
