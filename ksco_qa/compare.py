"""Work 결과와 최종 파이프라인 결과 비교 (설계 9.3, 10.1).

Work 결과는 별도 jsonl이 없다. 두 가지 입력을 받는다 (2026-10-21 사용자 확인):
- Work 결과 CSV/TSV: 열 id, AI 판정, 최종 판정, Regenerate Task, 원본 Task(백업), 검수 메모
- DataList.xlsx의 "DataList" 탭 R~V열 (R=AI 판정, S=최종 판정, T=Regenerate Task, U=원본 Task(백업), V=검수 메모), id는 B열
이 열들은 비교에만 쓰고 row_card(모델 입력)에는 절대 넣지 않는다 (CLAUDE.md M1 결정).
M5 비교 기준: Work "AI 판정" ↔ API final.status, Work "Regenerate Task" ↔
API revision.regenerated_task_new. Work "최종 판정"은 사용자가 지정하지 않아 비교하지 않는다.
"""

from __future__ import annotations

import csv
import io
import json
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Any

# Work 결과 필드 ← CSV 열 이름 / DataList 열 문자
WORK_RESULT_FIELDS = {
    "ai_verdict": ("AI 판정", "R"),
    "final_verdict": ("최종 판정", "S"),
    "regenerated_task": ("Regenerate Task", "T"),
    "original_task_backup": ("원본 Task(백업)", "U"),
    "review_memo": ("검수 메모", "V"),
}
WORK_RESULT_CSV_COLUMNS = ["id"] + [name for name, _ in WORK_RESULT_FIELDS.values()]


class WorkResultFormatError(Exception):
    """Work 결과 파일의 열 구조가 기대와 다름, id 누락·중복."""


class CompareError(Exception):
    """API 결과 JSONL이 없거나 비교에 필요한 구조가 잘못됨."""


@dataclass(frozen=True)
class WorkResultRow:
    id: str
    ai_verdict: str | None
    final_verdict: str | None
    regenerated_task: str | None
    original_task_backup: str | None
    review_memo: str | None
    source_row: int | None = None  # DataList에서 읽었을 때 엑셀 물리 행 번호


def _norm(value: Any) -> str:
    return unicodedata.normalize("NFC", "" if value is None else str(value)).strip()


def _value(value: Any) -> str | None:
    text = _norm(value)
    return text or None


def _check_ids(rows: list[WorkResultRow], source: str) -> dict[str, WorkResultRow]:
    out: dict[str, WorkResultRow] = {}
    for r in rows:
        if not r.id:
            raise WorkResultFormatError(f"{source}: id가 빈 행이 있다")
        if r.id in out:
            raise WorkResultFormatError(f"{source}: id 중복 {r.id}")
        out[r.id] = r
    return out


def load_work_result_csv(path: Path) -> dict[str, WorkResultRow]:
    """Work 결과 CSV/TSV를 id → 행으로 읽는다. 인코딩은 UTF-8(BOM 포함) → CP949 순서."""
    data = Path(path).read_bytes()
    for enc in ("utf-8-sig", "cp949"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise WorkResultFormatError(f"{Path(path).name}: UTF-8·CP949로 읽을 수 없다")
    delimiter = "\t" if Path(path).suffix.lower() == ".tsv" else ","
    reader = csv.DictReader(io.StringIO(text, newline=""), delimiter=delimiter)
    header = [_norm(h) for h in reader.fieldnames or []]
    if header != WORK_RESULT_CSV_COLUMNS:
        raise WorkResultFormatError(f"{Path(path).name}: 열이 {WORK_RESULT_CSV_COLUMNS}이어야 한다 (실제 {header})")
    rows = []
    for row_number, raw in enumerate(reader, start=2):
        if None in raw or any(raw.get(name) is None for name in WORK_RESULT_CSV_COLUMNS):
            raise WorkResultFormatError(f"{Path(path).name}:{row_number}행: 열 개수가 머리글과 다르다")
        row = {_norm(k): v for k, v in raw.items()}
        rows.append(WorkResultRow(_norm(row["id"]), *(_value(row[name]) for name, _ in WORK_RESULT_FIELDS.values())))
    return _check_ids(rows, Path(path).name)


def load_work_result_datalist(path: Path) -> dict[str, WorkResultRow]:
    """DataList.xlsx의 "DataList" 탭 R~V열을 id(B열) → 행으로 읽는다. 1행 머리글이 기대와 다르면 오류."""
    import openpyxl

    from ksco_qa.p0 import DATA_COLUMNS, SHEET_NAME, is_blank, norm_id

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:  # read_only는 파일 핸들을 잡고 있으므로 꼭 닫는다 (Windows 파일 잠금)
        return _read_work_result_sheet(wb, path, DATA_COLUMNS, SHEET_NAME, is_blank, norm_id)
    finally:
        wb.close()


def _read_work_result_sheet(wb: Any, path: Path, DATA_COLUMNS: str, SHEET_NAME: str, is_blank: Any,
                            norm_id: Any) -> dict[str, WorkResultRow]:
    if SHEET_NAME not in wb.sheetnames:
        raise WorkResultFormatError(f"{Path(path).name}: '{SHEET_NAME}' 탭이 없다")
    ws = wb[SHEET_NAME]
    rows_iter = ws.iter_rows(values_only=True)
    header = list(next(rows_iter, ()))
    col = {letter: ord(letter) - ord("A") for letter in "BCDEFGHIRSTUV"}
    wrong = [f"{letter}열 '{_norm(header[col[letter]]) if col[letter] < len(header) else ''}' (기대 '{name}')"
             for name, letter in WORK_RESULT_FIELDS.values()
             if col[letter] >= len(header) or _norm(header[col[letter]]) != name]
    if wrong:
        raise WorkResultFormatError(f"{Path(path).name}: Work 결과 열 머리글이 다르다: " + "; ".join(wrong))
    rows = []
    for r, values in enumerate(rows_iter, start=2):
        values = list(values) + [None] * (max(col.values()) + 1 - len(values))
        if all(is_blank(values[col[c]]) for c in DATA_COLUMNS):
            continue  # 데이터 행 = B~I 중 하나라도 값 (P0와 같은 기준)
        rows.append(WorkResultRow(norm_id(values[col["B"]]) or "",
                                  *(_value(values[col[letter]]) for _, letter in WORK_RESULT_FIELDS.values()),
                                  source_row=r))
    return _check_ids(rows, Path(path).name)


def load_work_result(path: Path) -> dict[str, WorkResultRow]:
    """compare --with 입력: .csv/.tsv면 Work 결과, .xlsx면 DataList R~V열."""
    suffix = Path(path).suffix.lower()
    if suffix in {".csv", ".tsv"}:
        return load_work_result_csv(path)
    if suffix == ".xlsx":
        return load_work_result_datalist(path)
    raise WorkResultFormatError(f"Work 결과는 .csv/.tsv 또는 DataList .xlsx여야 한다: {Path(path).name}")


def load_final_jsonl(results_dir: Path) -> list[dict[str, Any]]:
    """03_results/03_final.jsonl을 읽는다 (API 결과, 3장 원칙 8)."""
    path = Path(results_dir)
    if path.is_dir():
        path /= "03_final.jsonl"
    if not path.is_file():
        raise CompareError(f"API 최종 결과 파일이 없다: {path}")
    records = []
    seen_ids: set[str] = set()
    for line_number, line in enumerate(path.read_text(encoding="utf-8").splitlines(), start=1):
        if not line.strip():
            continue
        try:
            record = json.loads(line)
        except json.JSONDecodeError as error:
            raise CompareError(f"{path.name}:{line_number}: JSON 형식 오류: {error.msg}") from error
        if not isinstance(record, dict):
            raise CompareError(f"{path.name}:{line_number}: 레코드는 JSON 객체여야 한다")
        raw_id = record.get("id")
        if not isinstance(raw_id, str):
            raise CompareError(f"{path.name}:{line_number}: id는 문자열이어야 한다")
        row_id = _norm(raw_id)
        if not row_id:
            raise CompareError(f"{path.name}:{line_number}: id가 비어 있다")
        if row_id in seen_ids:
            raise CompareError(f"{path.name}:{line_number}: id 중복 {row_id}")
        if not isinstance(record.get("final"), dict) or not isinstance(record.get("revision"), dict):
            raise CompareError(f"{path.name}:{line_number}: final·revision 객체가 필요하다")
        if "status" not in record["final"] or "regenerated_task_new" not in record["revision"]:
            raise CompareError(f"{path.name}:{line_number}: final.status·revision.regenerated_task_new 필드가 필요하다")
        if record["final"]["status"] is not None and not isinstance(record["final"]["status"], str):
            raise CompareError(f"{path.name}:{line_number}: final.status는 문자열 또는 null이어야 한다")
        if (record["revision"]["regenerated_task_new"] is not None
                and not isinstance(record["revision"]["regenerated_task_new"], str)):
            raise CompareError(f"{path.name}:{line_number}: revision.regenerated_task_new는 문자열 또는 null이어야 한다")
        record["id"] = row_id
        seen_ids.add(row_id)
        records.append(record)
    return records


def _api_rows(api: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    rows: dict[str, dict[str, Any]] = {}
    for record in api:
        raw_id = record.get("id")
        if not isinstance(raw_id, str):
            raise CompareError("API 결과 id는 문자열이어야 한다")
        row_id = _norm(raw_id)
        if not row_id:
            raise CompareError("API 결과에 id가 비어 있는 행이 있다")
        if row_id in rows:
            raise CompareError(f"API 결과에 id 중복 {row_id}")
        if not isinstance(record.get("final"), dict) or not isinstance(record.get("revision"), dict):
            raise CompareError(f"API 결과 {row_id}: final·revision 객체가 필요하다")
        if "status" not in record["final"] or "regenerated_task_new" not in record["revision"]:
            raise CompareError(f"API 결과 {row_id}: final.status·revision.regenerated_task_new 필드가 필요하다")
        if record["final"]["status"] is not None and not isinstance(record["final"]["status"], str):
            raise CompareError(f"API 결과 {row_id}: final.status는 문자열 또는 null이어야 한다")
        task = record["revision"]["regenerated_task_new"]
        if task is not None and not isinstance(task, str):
            raise CompareError(f"API 결과 {row_id}: revision.regenerated_task_new는 문자열 또는 null이어야 한다")
        rows[row_id] = record
    return rows


def final_status_agreement(api: list[dict[str, Any]], work: dict[str, WorkResultRow]) -> dict[str, Any]:
    """Work AI 판정과 API final.status를 비교한다. 빈 값은 일치율 분모에서 제외한다."""
    api_rows = _api_rows(api)
    comparable = matching = 0
    missing_work, missing_api, mismatches = [], [], []
    for row_id, record in api_rows.items():
        api_status = _value(record["final"].get("status"))
        work_row = work.get(row_id)
        work_status = work_row.ai_verdict if work_row else None
        source_row = record.get("source_row")
        if work_row is None or work_status is None:
            missing_work.append({"id": row_id, "source_row": source_row})
            continue
        if api_status is None:
            missing_api.append({"id": row_id, "source_row": source_row, "work_ai_verdict": work_status})
            continue
        comparable += 1
        if api_status == work_status:
            matching += 1
        else:
            mismatches.append({
                "id": row_id, "source_row": source_row, "work_source_row": work_row.source_row,
                "api_final_status": api_status, "work_ai_verdict": work_status,
                "review_required": True,
            })
    return {
        "api_field": "final.status",
        "work_field": "AI 판정",
        "compared": comparable,
        "matching": matching,
        "mismatching": len(mismatches),
        "missing_work_value": missing_work,
        "missing_api_value": missing_api,
        "agreement_rate": round(matching / comparable, 4) if comparable else None,
        "mismatches": mismatches,
    }


def regenerated_task_diff(api: list[dict[str, Any]], work: dict[str, WorkResultRow]) -> dict[str, Any]:
    """Work Regenerate Task와 API revision.regenerated_task_new를 정규화 후 비교한다."""
    api_rows = _api_rows(api)
    compared = matching = 0
    mismatches = []
    for row_id, record in api_rows.items():
        work_row = work.get(row_id)
        if work_row is None:
            continue
        api_task = _value(record["revision"].get("regenerated_task_new"))
        work_task = work_row.regenerated_task
        if api_task is not None:
            api_task = api_task.replace("\r\n", "\n").replace("\r", "\n")
        if work_task is not None:
            work_task = work_task.replace("\r\n", "\n").replace("\r", "\n")
        compared += 1
        if api_task == work_task:
            matching += 1
        else:
            mismatches.append({
                "id": row_id,
                "source_row": record.get("source_row"),
                "work_source_row": work_row.source_row,
                "api_regenerated_task_new": api_task,
                "work_regenerate_task": work_task,
                "review_required": True,
            })
    return {
        "api_field": "revision.regenerated_task_new",
        "work_field": "Regenerate Task",
        "compared": compared,
        "matching": matching,
        "mismatching": len(mismatches),
        "agreement_rate": round(matching / compared, 4) if compared else None,
        "mismatches": mismatches,
    }


def compare_results(api_results: Path, work_result: Path) -> dict[str, Any]:
    """API final.status / regenerated task와 Work AI 판정 / regenerated task 비교 리포트."""
    api = load_final_jsonl(api_results)
    work = load_work_result(work_result)
    api_rows = _api_rows(api)
    api_ids, work_ids = set(api_rows), set(work)
    return {
        "schema_version": "ksco-qa/compare-1",
        "api_records": len(api),
        "work_records": len(work),
        "matched_ids": len(api_ids & work_ids),
        "api_only_ids": sorted(api_ids - work_ids),
        "work_only_ids": sorted(work_ids - api_ids),
        "final_status_agreement": final_status_agreement(api, work),
        "regenerated_task_comparison": regenerated_task_diff(api, work),
    }
