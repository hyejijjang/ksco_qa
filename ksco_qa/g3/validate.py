"""G3 검증 (설계 5.7-4, 9.1 G3 행 / 06_Result_Schema.md 최소 검증 1~10 / CLAUDE.md M1 결정 G3).

모든 검사는 실패 메시지 목록을 돌려준다 (빈 목록 = 통과). 하나라도 실패하면 run 상태는 FAILED_VALIDATION.
"""

from __future__ import annotations

import json
import re
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from ksco_qa.dictionaries import enum_values

TEMPLATE_MIN_ROWS = 10
MAX_MESSAGES = 20  # 검사마다 실패 메시지는 이만큼만 적고 나머지는 개수로


def _cap(errors: list[str]) -> list[str]:
    return errors if len(errors) <= MAX_MESSAGES else errors[:MAX_MESSAGES] + [f"… 외 {len(errors) - MAX_MESSAGES}건"]


def read_final_jsonl(path: Path) -> tuple[list[dict[str, Any]], list[str]]:
    """물리 줄마다 JSON 파싱 (00 8장). (레코드, 파싱 오류)."""
    records, errors = [], []
    for i, line in enumerate(Path(path).read_text(encoding="utf-8").split("\n"), start=1):
        if not line.strip():
            continue
        try:
            records.append(json.loads(line))
        except json.JSONDecodeError as e:
            errors.append(f"{i}번째 줄 JSON 파싱 실패: {e.msg}")
    return records, errors


def _xlsx_rows(path: Path) -> list[list[Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:  # read_only는 파일 핸들을 잡고 있으므로 꼭 닫는다 (Windows 파일 잠금)
        return [list(r) for r in wb.active.iter_rows(values_only=True)]
    finally:
        wb.close()


# --- 06 최소 검증 1~10 ---


def check_row_counts(datalist_rows: int, gate_files: dict[str, Path], records: list[dict[str, Any]],
                     mapping_audit: Path | None = None) -> list[str]:
    """1. 원본 데이터 행 수 = 모든 Gate 파일 행 수 = JSONL 레코드 수."""
    errors = []
    if len(records) != datalist_rows:
        errors.append(f"03_final.jsonl {len(records)}행 ≠ 원본 {datalist_rows}행")
    for name, path in gate_files.items():
        n = len(_xlsx_rows(path)) - 1
        if n != datalist_rows:
            errors.append(f"{name} {n}행 ≠ 원본 {datalist_rows}행")
    if mapping_audit is not None:
        with open(mapping_audit, encoding="utf-8-sig") as f:
            n = sum(1 for _ in f) - 1
        if n != datalist_rows:
            errors.append(f"00_mapping_audit.csv {n}행 ≠ 원본 {datalist_rows}행")
    return errors


def check_row_identity(row_cards: list[dict[str, Any]], records: list[dict[str, Any]]) -> list[str]:
    """2. source_row와 id가 원본 행에 대응, 정렬·삭제·삽입 없음."""
    expected = [(c["source_row"], c["id"]) for c in row_cards]
    actual = [(r.get("source_row"), r.get("id")) for r in records]
    if expected == actual:
        return []
    errors = [f"{i + 1}번째: 원본 {e} ≠ JSONL {a}" for i, (e, a) in enumerate(zip(expected, actual)) if e != a]
    if len(expected) != len(actual):
        errors.append(f"행 수 다름: 원본 {len(expected)}, JSONL {len(actual)}")
    return _cap(errors)


def _same(a: Any, b: Any) -> bool:
    if a is None or a == "":
        return b is None or b == ""
    if isinstance(a, bool) or isinstance(b, bool):
        return a == b
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return float(a) == float(b)
    return str(a) == str(b)


def check_source_cells(source_cells: dict[int, dict[str, Any]], records: list[dict[str, Any]],
                       gate_files: dict[str, Path]) -> list[str]:
    """3. 각 JSONL source.cells의 원본 셀과 원본(P0 source_cells), 최종 Excel 기존 영역 대조."""
    errors = []
    for r in records:
        original = (source_cells.get(r["source_row"]) or {}).get("cells")
        if original != r["source"]["cells"]:
            errors.append(f"{r['source_row']}행 source.cells가 원본과 다르다")
    for name, path in gate_files.items():
        rows = _xlsx_rows(path)
        for r, row in zip(records, rows[1:]):
            letters = sorted(r["source"]["cells"], key=lambda c: (len(c), c))
            for i, letter in enumerate(letters):
                if not _same(r["source"]["cells"][letter]["value"], row[i] if i < len(row) else None):
                    errors.append(f"{name} {r['source_row']}행 {letter}열이 원본 값과 다르다")
                    break
    return _cap(errors)


def check_match_method(records: list[dict[str, Any]]) -> list[str]:
    """4. D_F_EXACT_FILENAME 외의 매핑 방법이 MATCHED로 기록되지 않았는가."""
    return _cap([f"{r['source_row']}행 MATCHED인데 match_method={d['match_method']!r}"
                 for r in records for d in [r["source_assets"]["deliverable_files"]]
                 if d["match_status"] == "MATCHED" and d["match_method"] != "D_F_EXACT_FILENAME"])


def check_regenerate_task(records: list[dict[str, Any]]) -> list[str]:
    """5. G0 REGENERATE 행에 regenerated_task_new가 있고 effective_task와 정확히 일치하는가."""
    errors = []
    for r in records:
        rev = r["revision"]
        if rev["action"] != "REGENERATE":
            continue
        if not rev["regenerated_task_new"]:
            errors.append(f"{r['source_row']}행 REGENERATE인데 regenerated_task_new 없음")
        elif rev["regenerated_task_new"] != rev["effective_task"]:
            errors.append(f"{r['source_row']}행 regenerated_task_new ≠ effective_task")
    return _cap(errors)


def check_duplicate_consistency(records: list[dict[str, Any]]) -> list[str]:
    """6. DUPLICATE면 대표 행 존재(대표 자신은 DUPLICATE 아님), verify SKIPPED, final FAIL."""
    by_row = {r["source_row"]: r for r in records}
    errors = []
    for r in records:
        dup = r["duplicate"]
        if dup["status"] != "DUPLICATE":
            continue
        rep = by_row.get(dup["representative_row"])
        if rep is None or dup["representative_id"] != rep["id"]:
            errors.append(f"{r['source_row']}행 DUPLICATE인데 대표 행·id가 원본에 없거나 맞지 않음")
        elif rep["duplicate"]["status"] == "DUPLICATE":
            errors.append(f"{r['source_row']}행의 대표 {rep['source_row']}행도 DUPLICATE")
        if r["verify"]["status"] != "SKIPPED":
            errors.append(f"{r['source_row']}행 DUPLICATE인데 verify.status={r['verify']['status']}")
        if r["final"]["status"] != "FAIL":
            errors.append(f"{r['source_row']}행 DUPLICATE인데 final.status={r['final']['status']}")
    return _cap(errors)


def check_search_completed(records: list[dict[str, Any]]) -> list[str]:
    """7. G1 미평가 후보가 있는데 search_completed=true라고 표기하지 않았는가."""
    return _cap([f"{r['source_row']}행 미판정 후보 {r['duplicate']['unresolved_candidate_count']}개인데 search_completed=true"
                 for r in records if r["duplicate"]["unresolved_candidate_count"] and r["duplicate"]["search_completed"]])


def check_reference_not_auto_review(records: list[dict[str, Any]]) -> list[str]:
    """8. Reference 부재(GENERATE_REFERENCE)를 이유로 자동 REVIEW하지 않았는가."""
    ref_codes = {"REFERENCE_UNRESOLVED", "REFERENCE_HOLD", "REFERENCE_CONFLICT"}
    return _cap([f"{r['source_row']}행 Reference 부재인데 {sorted(ref_codes & set(r['final']['reason_codes']))}로 REVIEW"
                 for r in records
                 if r["source_assets"]["reference_files"]["mode"] == "GENERATE_REFERENCE"
                 and ref_codes & set(r["final"]["reason_codes"])])


def check_pass_conditions(records: list[dict[str, Any]]) -> list[str]:
    """9. final PASS면 G0 채택, G1 완료·UNIQUE, G2 PASS, 핵심 자료 확인(MATCHED·inspected), readiness≠HOLD."""
    errors = []
    for r in records:
        if r["final"]["status"] != "PASS":
            continue
        d = r["source_assets"]["deliverable_files"]
        problems = [p for p, bad in (
            ("G0 미채택", r["revision"]["action"] not in ("KEEP", "REGENERATE")),
            ("G1 UNIQUE·완료 아님", r["duplicate"]["status"] != "UNIQUE" or not r["duplicate"]["search_completed"]),
            ("G2 PASS 아님", r["verify"]["status"] != "PASS"),
            ("자료 미확인", d["match_status"] != "MATCHED" or not d["inspected"]),
            ("Reference HOLD", r["verify"]["reference_readiness"] == "HOLD"),
        ) if bad]
        if problems:
            errors.append(f"{r['source_row']}행 PASS인데 {', '.join(problems)}")
    return _cap(errors)


def check_summary_counts(records: list[dict[str, Any]], summary: dict[str, Any]) -> list[str]:
    """10. 집계(run_report summary — QA_SUMMARY 대신)가 행 수·후보쌍 수와 맞는가."""
    from ksco_qa.g3.export import summarize

    fresh = summarize(records)
    errors = [f"summary.{k} {summary.get(k)} ≠ 레코드 집계 {v}" for k, v in fresh.items() if summary.get(k) != v]
    for key in ("g0_action", "g1_status", "g2_status", "final_status"):
        total = sum((summary.get(key) or {}).values())
        if total != len(records):
            errors.append(f"summary.{key} 합계 {total} ≠ 행 수 {len(records)}")
    return errors


# --- 추가 검사 (설계 9.1 G3 행) ---

ENUM_FIELDS = {
    ("revision", "action"): "revision_action",
    ("duplicate", "status"): "duplicate_status",
    ("duplicate", "duplicate_type"): "duplicate_type",
    ("duplicate", "overlap_status"): "overlap_status",
    ("verify", "status"): "verify_status",
    ("verify", "b_coverage_status"): "b_coverage_status",
    ("verify", "reference_mode"): "reference_mode",
    ("verify", "reference_action"): "reference_action",
    ("verify", "reference_readiness"): "reference_readiness",
    ("final", "status"): "final_status",
    ("final", "next_action"): "next_action",
}
# 값이 없을 수 있는 enum (규칙상 SKIPPED·미평가 행에서 비는 필드)
ENUM_NULLABLE = {("verify", "b_coverage_status"), ("verify", "reference_action"), ("verify", "reference_readiness"),
                 ("duplicate", "duplicate_type")}


def check_enums(records: list[dict[str, Any]]) -> list[str]:
    """enum 필드가 enums.yaml 허용값인가 (DESIGNABLE 같은 값 차단, 7.5). overlap_status는 null 금지."""
    errors = []
    codes = set(enum_values("reason_code"))
    decided = set(enum_values("decided_by"))
    statuses = set(enum_values("check_status_final"))
    for r in records:
        row = r["source_row"]
        for (gate, fld), enum in ENUM_FIELDS.items():
            value = r[gate][fld]
            if value is None and (gate, fld) in ENUM_NULLABLE:
                continue
            if value not in enum_values(enum):
                errors.append(f"{row}행 {gate}.{fld}={value!r} (허용: {enum})")
        for gate in ("revision", "verify", "final"):
            bad = set(r[gate].get("reason_codes") or []) - codes
            if bad:
                errors.append(f"{row}행 {gate}.reason_codes에 허용되지 않은 값 {sorted(bad)}")
        for gate in ("revision", "duplicate", "verify", "final"):
            if r[gate].get("decided_by") not in decided:
                errors.append(f"{row}행 {gate}.decided_by={r[gate].get('decided_by')!r}")
        for gate, key in (("revision", "g0_self_check"), ("verify", "checks")):
            for name, check in r[gate][key].items():
                if check["status"] not in statuses:
                    errors.append(f"{row}행 {gate}.{key}.{name}.status={check['status']!r}")
    return _cap(errors)


# 규칙상 모든 행에서 비어 있을 수 있는 필드 (예: REGENERATE 행이 없으면 regenerated_task_new는 전부 null)
CONDITIONALLY_NULL = {
    "revision.regenerated_task_new", "revision.effective_task", "revision.evidence", "revision.expected_reference_spec",
    "duplicate.representative_row", "duplicate.representative_id", "duplicate.group", "duplicate.duplicate_type",
    "verify.skip_reason", "verify.proposed_fix", "verify.reference_conflict", "verify.reference_conflict_note",
    "verify.b_coverage_status", "verify.reference_action", "verify.reference_readiness", "verify.reason",
    "source_assets.deliverable_files.match_method", "source_assets.deliverable_files.matched_locator",
    "final.rule",
}


def _leaves(obj: Any, prefix: str = "") -> dict[str, Any]:
    out: dict[str, Any] = {}
    if isinstance(obj, dict) and obj:
        for k, v in obj.items():
            out.update(_leaves(v, f"{prefix}.{k}" if prefix else k))
    else:
        out[prefix] = obj
    return out


def check_all_null_fields(records: list[dict[str, Any]]) -> list[str]:
    """전 행 null인 필드 = 실패 (교훈 #9). source.cells 원본과 규칙상 비어도 되는 필드는 제외."""
    if not records:
        return []
    seen: dict[str, bool] = defaultdict(bool)
    for r in records:
        body = {k: v for k, v in r.items() if k != "source"}
        for path, value in _leaves(body).items():
            seen[path] |= value is not None
    return [f"{path}가 전 행 null" for path, any_value in sorted(seen.items())
            if not any_value and path not in CONDITIONALLY_NULL and not path.startswith(
                ("revision.g0_self_check.", "verify.checks.", "verify.blueprints."))]


_QUOTED_RE = re.compile(r"'[^']*'|\"[^\"]*\"|‘[^’]*’|“[^”]*”|「[^」]*」|『[^』]*』|\([^)]*\)|\[[^\]]*\]|<[^>]*>")


def normalize_template(text: str, row_terms: list[str]) -> str:
    """숫자, 해당 행의 job·job_leaf·deliverable_title, 따옴표·괄호 안 값, 공백을 지운다 (M1 결정 22)."""
    for term in sorted((t for t in row_terms if t), key=len, reverse=True):
        text = text.replace(term, "")
    text = _QUOTED_RE.sub("", text)
    text = re.sub(r"\d+", "", text)
    return re.sub(r"\s+", "", text)


def model_texts(r: dict[str, Any]) -> list[tuple[str, str]]:
    """템플릿 감지 대상 (모델이 쓴 글만): G0 reason·issues, G1b evidence, G2 reason·issues·proposed_fix."""
    out = []
    rev, ver = r["revision"], r["verify"]
    if rev["decided_by"] == "MODEL":
        if rev["reason"]:
            out.append(("G0 reason", rev["reason"]))
        out += [(f"G0 issues.{k}", c["conclusion"]) for k, c in rev["g0_self_check"].items()
                if c["status"] in ("REVIEW", "FAIL") and c["conclusion"]]
    out += [("G1b evidence", p["evidence"]) for p in r["duplicate"]["candidate_pairs"] if p.get("evidence")]
    if ver["decided_by"] == "MODEL":
        if ver["reason"]:
            out.append(("G2 reason", ver["reason"]))
        if ver["proposed_fix"]:
            out.append(("G2 proposed_fix", ver["proposed_fix"]))
        out += [(f"G2 issues.{k}", c["conclusion"]) for k, c in ver["checks"].items()
                if c["status"] in ("REVIEW", "FAIL") and c["conclusion"]]
    return out


def check_template_sentences(records: list[dict[str, Any]], row_cards: list[dict[str, Any]],
                             min_rows: int = TEMPLATE_MIN_ROWS) -> list[str]:
    """같은 문장(정규화 후 완전 일치)이 min_rows행 이상이면 템플릿 판정으로 실패."""
    cards = {c["source_row"]: c for c in row_cards}
    rows_by_text: dict[tuple[str, str], set[int]] = defaultdict(set)
    for r in records:
        card = cards.get(r["source_row"], {})
        terms = [card.get("job") or "", card.get("job_leaf") or "", card.get("deliverable_title") or ""]
        for field, text in model_texts(r):
            key = normalize_template(text, terms)
            if key:
                rows_by_text[(field.split(".")[0], key)].add(r["source_row"])
    return [f"{field} 같은 문장이 {len(rows)}행에 반복 (예: {sorted(rows)[:5]}행)"
            for (field, _), rows in rows_by_text.items() if len(rows) >= min_rows]


def validate_all(results: dict[str, Path], inputs: Any, summary: dict[str, Any]) -> dict[str, list[str]]:
    """06 최소 검증 1~10 + enum·전 행 null·템플릿 감지. 검사 이름 → 실패 메시지 목록."""
    from ksco_qa.g3.export import FINAL_JSONL, GATE_FILES

    records, parse_errors = read_final_jsonl(results[FINAL_JSONL])
    gate_files = {name: results[name] for name in GATE_FILES}
    mapping_audit = results.get("00_mapping_audit.csv")
    return {
        "00_jsonl_parse": parse_errors,
        "01_row_counts": check_row_counts(len(inputs.row_cards), gate_files, records,
                                          mapping_audit if mapping_audit and Path(mapping_audit).is_file() else None),
        "02_row_identity": check_row_identity(inputs.row_cards, records),
        "03_source_cells": check_source_cells(inputs.source_cells, records, gate_files),
        "04_match_method": check_match_method(records),
        "05_regenerate_task": check_regenerate_task(records),
        "06_duplicate_consistency": check_duplicate_consistency(records),
        "07_search_completed": check_search_completed(records),
        "08_reference_not_auto_review": check_reference_not_auto_review(records),
        "09_pass_conditions": check_pass_conditions(records),
        "10_summary_counts": check_summary_counts(records, summary),
        "enums": check_enums(records),
        "all_null_fields": check_all_null_fields(records),
        "template_sentences": check_template_sentences(records, inputs.row_cards),
    }
