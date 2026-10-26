"""같은 표본을 두 번 돌린 run의 행별 비교표 (2026-10-26, M2 게이트 사람 검토용).

이전 run과 새 run의 G0 action·차원별 R/F·재생성 Task, G1·G2·최종 상태를 행마다 나란히 놓고,
Work 결과(사람 최종 판정·검수 메모)를 붙인다. 모델을 부르지 않고 run 폴더 기록만 읽는다.
Work 결과 열은 비교에만 쓰고 모델 입력에는 넣지 않는다 (CLAUDE.md M1 결정).
"""

from __future__ import annotations

import json
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import Workbook

from ksco_qa.compare import CompareError, WorkResultRow, load_work_result
from ksco_qa.dictionaries import g0_check_keys

ROW_HEADERS = [
    "source_row", "id", "job_leaf", "사람 AI 판정", "사람 최종 판정", "검수 메모", "Work Regenerate Task",
    "이전 action", "새 action", "action 변경", "이전 R/F 수", "새 R/F 수", "이전 R/F 차원", "새 R/F 차원",
    "이전 재판정", "새 재판정", "원본 Task", "이전 재생성 Task", "새 재생성 Task", "새 reason",
    "이전 G1", "새 G1", "이전 G2", "새 G2", "이전 최종", "새 최종",
]


def _jsonl(path: Path) -> list[dict[str, Any]]:
    if not path.is_file():
        return []
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def _by_row(records: list[dict[str, Any]]) -> dict[int, dict[str, Any]]:
    """추가 전용 기록이므로 같은 행은 마지막 줄이 이긴다."""
    return {int(r["source_row"]): r for r in records}


def _final_records(run_dir: Path) -> dict[int, dict[str, Any]]:
    """run 폴더의 03_final.jsonl, 없으면 03_results/03_final.jsonl 중 이 run이 내보낸 것."""
    run_id = run_dir.name
    for path in (run_dir / "03_final.jsonl", run_dir.parent.parent / "03_final.jsonl"):
        records = _jsonl(path)
        if records and f"run_id={run_id};" in str(records[0].get("final", {}).get("execution_notes")):
            return _by_row(records)
    return {}


def load_run(run_dir: Path) -> dict[str, Any]:
    run_dir = Path(run_dir)
    if not (run_dir / "g0" / "records.jsonl").is_file():
        raise CompareError(f"G0 기록이 없다: {run_dir / 'g0' / 'records.jsonl'}")
    report = run_dir / "run_report.json"
    return {
        "run_id": run_dir.name,
        "cards": _by_row(_jsonl(run_dir / "row_cards.jsonl")),
        "g0": _by_row(_jsonl(run_dir / "g0" / "records.jsonl")),
        "g1": _by_row(_jsonl(run_dir / "g1" / "rows.jsonl")),
        "g2": _by_row(_jsonl(run_dir / "g2" / "records.jsonl")),
        "final": _final_records(run_dir),
        "calls": _jsonl(run_dir / "calls.jsonl"),
        "report": json.loads(report.read_text(encoding="utf-8")) if report.is_file() else {},
    }


def _flagged(g0_record: dict[str, Any] | None) -> list[str]:
    checks = ((g0_record or {}).get("output") or {}).get("checks") or {}
    return [f"{key}:{checks[key]}" for key in g0_check_keys() if checks.get(key) in ("R", "F")]


def _g0_value(g0_record: dict[str, Any] | None, key: str) -> Any:
    return ((g0_record or {}).get("output") or {}).get(key)


def _status(run: dict[str, Any], stage: str, row: int) -> str | None:
    if stage == "final":
        return (run["final"].get(row) or {}).get("final", {}).get("status")
    return (run[stage].get(row) or {}).get("status")


def build_run_comparison(old_dir: Path, new_dir: Path, work: dict[str, WorkResultRow]) -> dict[str, Any]:
    old, new = load_run(old_dir), load_run(new_dir)
    rows = []
    for row in sorted(new["g0"]):
        n, o = new["g0"][row], old["g0"].get(row)
        card = new["cards"].get(row, {})
        w = work.get(n["id"])
        old_flags, new_flags = _flagged(o), _flagged(n)
        rows.append({
            "source_row": row, "id": n["id"], "job_leaf": card.get("job_leaf"),
            "사람 AI 판정": w.ai_verdict if w else None, "사람 최종 판정": w.final_verdict if w else None,
            "검수 메모": w.review_memo if w else None, "Work Regenerate Task": w.regenerated_task if w else None,
            "이전 action": (o or {}).get("action"), "새 action": n.get("action"),
            "action 변경": (o or {}).get("action") != n.get("action"),
            "이전 R/F 수": len(old_flags) if o else None, "새 R/F 수": len(new_flags),
            "이전 R/F 차원": ", ".join(old_flags), "새 R/F 차원": ", ".join(new_flags),
            "이전 재판정": (o or {}).get("escalated"), "새 재판정": n.get("escalated"),
            "원본 Task": card.get("original_task"),
            "이전 재생성 Task": _g0_value(o, "regenerated_task_new"), "새 재생성 Task": _g0_value(n, "regenerated_task_new"),
            "새 reason": _g0_value(n, "reason") or n.get("rule_reason"),
            "이전 G1": _status(old, "g1", row), "새 G1": _status(new, "g1", row),
            "이전 G2": _status(old, "g2", row), "새 G2": _status(new, "g2", row),
            "이전 최종": _status(old, "final", row), "새 최종": _status(new, "final", row),
        })

    dimensions = []
    for key in g0_check_keys():
        counts = {}
        for label, run in (("이전", old), ("새", new)):
            values = Counter(((r.get("output") or {}).get("checks") or {}).get(key) for r in run["g0"].values())
            counts[f"{label} R"], counts[f"{label} F"] = values["R"], values["F"]
        dimensions.append({"차원": key, **counts})

    def counts(column: str) -> dict[str, int]:
        return dict(Counter(r[column] for r in rows if r[column] is not None))

    human_pass = [r for r in rows if r["사람 최종 판정"] == "PASS"]
    judged = [r for r in rows if r["사람 최종 판정"]]
    classes = sorted({r["사람 최종 판정"] for r in judged})
    class_recall = {}  # 전체 일치율은 다수 클래스(REVIEW)에 끌려가므로 클래스별로 본다 (2026-10-26)
    for label, column, run in (("이전", "이전 최종", old), ("새", "새 최종", new)):
        if not run["final"]:  # 03_final.jsonl이 다른 run에 덮어써졌으면 비교할 수 없다
            class_recall[f"사람 판정 대비 최종 ({label})"] = "최종 기록 없음 (03_final.jsonl이 이 run 것이 아님)"
            continue
        for c in classes:
            members = [r for r in judged if r["사람 최종 판정"] == c]
            class_recall[f"사람 {c} 중 우리 {c} ({label})"] = f"{sum(r[column] == c for r in members)}/{len(members)}"
        hits = sum(r[column] == r["사람 최종 판정"] for r in judged)
        class_recall[f"전체 일치 ({label})"] = f"{hits}/{len(judged)} ({hits / len(judged):.0%})" if judged else None
    for c in classes:
        n = sum(r["사람 최종 판정"] == c for r in judged)
        class_recall[f"기준선: 전부 {c}로 답하면 전체 일치"] = f"{n}/{len(judged)} ({n / len(judged):.0%})"
    summary = {
        "이전 run": old["run_id"], "새 run": new["run_id"], "행 수": len(rows),
        "이전 action": counts("이전 action"), "새 action": counts("새 action"),
        "action 바뀐 행": sum(r["action 변경"] for r in rows),
        "Work Regenerate Task 있는 행": sum(bool(r["Work Regenerate Task"]) for r in rows),
        "사람 최종 PASS 행": len(human_pass),
        "사람 최종 PASS 중 REGENERATE (이전)": sum(r["이전 action"] == "REGENERATE" for r in human_pass),
        "사람 최종 PASS 중 REGENERATE (새)": sum(r["새 action"] == "REGENERATE" for r in human_pass),
        "G0 재판정 호출 (이전)": sum(c.get("stage") == "G0" and c.get("attempt_kind") == "escalate" for c in old["calls"]),
        "G0 재판정 호출 (새)": sum(c.get("stage") == "G0" and c.get("attempt_kind") == "escalate" for c in new["calls"]),
        "이전 G1": counts("이전 G1"), "새 G1": counts("새 G1"),
        "이전 G2": counts("이전 G2"), "새 G2": counts("새 G2"),
        "이전 최종": counts("이전 최종"), "새 최종": counts("새 최종"),
        **class_recall,
        "이전 비용 USD": old["report"].get("usage", {}).get("total", {}).get("cost_usd"),
        "새 비용 USD": new["report"].get("usage", {}).get("total", {}).get("cost_usd"),
        "새 템플릿 감지": new["report"].get("validation", {}).get("template_sentences"),
    }
    return {"rows": rows, "dimensions": dimensions, "summary": summary}


def _cell(value: Any) -> Any:
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return value


def write_run_comparison_xlsx(path: Path, comparison: dict[str, Any]) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "행별"
    sheet.append(ROW_HEADERS)
    for row in comparison["rows"]:
        sheet.append([_cell(row[h]) for h in ROW_HEADERS])
    dims = workbook.create_sheet("차원별 RF")
    headers = ["차원", "이전 R", "이전 F", "새 R", "새 F"]
    dims.append(headers)
    for row in comparison["dimensions"]:
        dims.append([row[h] for h in headers])
    summary = workbook.create_sheet("요약")
    for key, value in comparison["summary"].items():
        summary.append([key, _cell(value)])
    path = Path(path)
    workbook.save(path)
    return path


def compare_runs(task_folder: Path, old_run: str, new_run: str, work_result: Path, out: Path | None = None) -> Path:
    work_root = Path(task_folder) / "03_results" / "_work"
    old_dir, new_dir = work_root / old_run, work_root / new_run
    for run_dir in (old_dir, new_dir):
        if not run_dir.is_dir():
            raise CompareError(f"run 폴더가 없다: {run_dir}")
    comparison = build_run_comparison(old_dir, new_dir, load_work_result(work_result))
    return write_run_comparison_xlsx(out or new_dir / f"compare_vs_{old_run}.xlsx", comparison)
