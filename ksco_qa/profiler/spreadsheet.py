"""XLSX·XLS·ODS 프로필 (설계 부록 B XLSX·XLS·ODS 행).

시트 목록, 시트별 사용 범위, 머리글 영역(병합 셀 포함), 수식·합계 열, 앞·중간·끝 샘플.
머리글이 여러 행인 통계표가 많아 위 3행을 머리글 영역으로 고정 표시한다.
수식은 캐시값과 원문을 모두 확인한다 (XLS는 수식 원문 확인 불가).
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from ksco_qa.profiler.base import TOTAL_LABEL_RE, ProfileError, ProfileParts, is_blank, ranges, row_text
from ksco_qa.profiler.doc_model import FIRST_COL_MAX, DocTable, labels_line, table_layout

HEADER_ROWS = 3
SAMPLE = 3
MAX_MERGES = 8
LARGE_FILE = 20 * 1024 * 1024
ODS_MAX_ROWS = 200_000


@dataclass
class SheetData:
    name: str
    rows: list[list[Any]]  # 값(캐시값), 1행부터. 끝의 빈 행은 잘라 냄
    formula_counts: dict[str, int] | None = None  # 열 문자 → 수식 셀 수 (None = 확인 불가)
    merges: list[str] | None = None  # 머리글 영역의 병합 범위 (None = 확인 불가)
    notes: list[str] = field(default_factory=list)


def col_letter(index: int) -> str:
    """0 → A."""
    s, n = "", index + 1
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def _trim(rows: list[list[Any]]) -> list[list[Any]]:
    while rows and all(is_blank(v) for v in rows[-1]):
        rows.pop()
    return rows


def _load_xlsx(path: Path) -> list[SheetData]:
    import openpyxl

    large = path.stat().st_size > LARGE_FILE
    wb_v = openpyxl.load_workbook(path, data_only=True, read_only=large)
    wb_f = openpyxl.load_workbook(path, data_only=False, read_only=large)
    try:  # read_only는 파일 핸들을 잡고 있으므로 꼭 닫는다 (Windows 파일 잠금)
        return _read_xlsx_sheets(wb_v, wb_f, large)
    finally:
        wb_v.close()
        wb_f.close()


def _read_xlsx_sheets(wb_v: Any, wb_f: Any, large: bool) -> list[SheetData]:
    sheets = []
    for name in wb_v.sheetnames:
        ws_v, ws_f = wb_v[name], wb_f[name]
        rows = _trim([list(r) for r in ws_v.iter_rows(values_only=True)])
        counts: dict[str, int] = {}
        for row in ws_f.iter_rows():
            for cell in row:
                if getattr(cell, "data_type", None) == "f":
                    letter = col_letter(cell.column - 1)
                    counts[letter] = counts.get(letter, 0) + 1
        merges = None
        notes = []
        if large:
            notes.append("대용량이라 병합 셀 미확인")
        else:
            merges = [str(r) for r in ws_f.merged_cells.ranges if r.min_row <= HEADER_ROWS]
        sheets.append(SheetData(name, rows, counts, merges, notes))
    return sheets


def _load_xls(path: Path) -> list[SheetData]:
    import xlrd

    try:
        book = xlrd.open_workbook(path, formatting_info=True)
        with_merges = True
    except (NotImplementedError, xlrd.XLRDError):
        book = xlrd.open_workbook(path)
        with_merges = False
    sheets = []
    for sh in book.sheets():
        rows = _trim([sh.row_values(r) for r in range(sh.nrows)])
        merges = None
        if with_merges:
            merges = [f"{col_letter(c1)}{r1 + 1}:{col_letter(c2 - 1)}{r2}" for r1, r2, c1, c2 in sh.merged_cells
                      if r1 < HEADER_ROWS]
        sheets.append(SheetData(sh.name, rows, None, merges, ["XLS 형식: 수식 원문 확인 불가"]))
    return sheets


def _load_ods(path: Path) -> list[SheetData]:
    from odf import teletype
    from odf.opendocument import load
    from odf.table import Table, TableRow

    doc = load(str(path))
    sheets = []
    for table in doc.spreadsheet.getElementsByType(Table):
        rows: list[list[Any]] = []
        counts: dict[str, int] = {}
        merges: list[str] = []
        for tr in table.getElementsByType(TableRow):
            row_repeat = int(tr.getAttribute("numberrowsrepeated") or 1)
            cells: list[Any] = []
            for tc in tr.childNodes:
                if tc.qname[1] not in ("table-cell", "covered-table-cell"):
                    continue
                repeat = int(tc.getAttribute("numbercolumnsrepeated") or 1) if tc.qname[1] == "table-cell" else 1
                text = teletype.extractText(tc) if tc.qname[1] == "table-cell" else ""
                formula = tc.getAttribute("formula") if tc.qname[1] == "table-cell" else None
                span = tc.getAttribute("numbercolumnsspanned") if tc.qname[1] == "table-cell" else None
                if span and int(span) > 1 and len(rows) < HEADER_ROWS:
                    c0 = len(cells)
                    merges.append(f"{col_letter(c0)}{len(rows) + 1}:{col_letter(c0 + int(span) - 1)}{len(rows) + 1}")
                if formula:
                    letter = col_letter(len(cells))
                    counts[letter] = counts.get(letter, 0) + repeat
                cells.extend([text or None] * min(repeat, 1024))
            while cells and is_blank(cells[-1]):
                cells.pop()
            if not cells and row_repeat > 1:
                rows.extend([[]] * min(row_repeat, 10))  # 대량 반복 빈 행은 끝에서 잘린다
            else:
                rows.extend([cells] * min(row_repeat, ODS_MAX_ROWS))
            if len(rows) > ODS_MAX_ROWS:
                break
        sheets.append(SheetData(table.getAttribute("name"), _trim(rows), counts, merges))
    return sheets


_RANGE_RE = re.compile(r"^([A-Z]+)(\d+):([A-Z]+)(\d+)$")


def _col_index(letters: str) -> int:
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n - 1


def _row_labels_line(s: SheetData) -> str | None:
    """행 라벨 (PDF·문서형 표와 같은 규칙): 위에서부터 제목 행을 건너뛰고 머리글 행을 찾아 그 아래 행에 적용."""
    if not s.rows:
        return None
    n_cols = max(len(r) for r in s.rows)
    full_span = set()
    for rng in s.merges or []:
        m = _RANGE_RE.match(rng)
        if m and m.group(2) == m.group(4) and _col_index(m.group(3)) - _col_index(m.group(1)) + 1 >= n_cols:
            full_span.add(int(m.group(2)) - 1)
    rows = [["" if v is None else str(v) for v in r] for r in s.rows]
    table = DocTable(rows, 0, False, full_span)
    layout = table_layout(table)
    line = labels_line(table, FIRST_COL_MAX, indent="  ")
    return f"{line} (머리글 {layout.header + 1}행 기준)" if line else None


LOADERS = {"xlsx": _load_xlsx, "xlsm": _load_xlsx, "xls": _load_xls, "ods": _load_ods}


def profile(path: Path) -> ProfileParts:
    ext = path.suffix.lower().lstrip(".")
    sheets = LOADERS[ext](path)
    if not sheets:
        raise ProfileError("시트가 없다")
    fixed = [f"{ext.upper()} · 시트 {len(sheets)}개: {', '.join(s.name for s in sheets)}"]
    locations, head, mid, tail = [], [], [], []
    for s in sheets:
        n_rows = len(s.rows)
        n_cols = max((len(r) for r in s.rows), default=0)
        used = f"A1:{col_letter(n_cols - 1)}{n_rows}" if n_rows and n_cols else "빈 시트"
        fixed.append(f"[{s.name}] 사용 범위 {used} ({n_rows}행 × {n_cols}열)")
        for i, r in enumerate(s.rows[:HEADER_ROWS]):
            fixed.append(f"  머리글 {i + 1}행: {row_text(r)}")
        if s.merges is not None:
            shown = s.merges[:MAX_MERGES]
            more = f" 외 {len(s.merges) - len(shown)}개" if len(s.merges) > len(shown) else ""
            fixed.append(f"  머리글 병합 셀: {', '.join(shown) + more if shown else '없음'}")
        if s.formula_counts is not None:
            f_cols = ", ".join(f"{c}({n})" for c, n in sorted(s.formula_counts.items()))
            fixed.append(f"  수식 열: {f_cols or '없음'}")
        totals = [i + 1 for i, r in enumerate(s.rows)
                  if r and TOTAL_LABEL_RE.match(str(next((c for c in r if not is_blank(c)), "")))]
        fixed.append(f"  합계·소계 행: {ranges(totals[:20]) + '행' if totals else '없음'}")
        label = _row_labels_line(s)
        if label:
            fixed.append(label)
        fixed.extend(f"  ({note})" for note in s.notes)

        body = list(range(HEADER_ROWS, n_rows))
        h = body[:SAMPLE]
        t = body[-SAMPLE:] if len(body) > SAMPLE else []
        t = [i for i in t if i not in h]
        m_start = len(body) // 2
        m = [i for i in body[m_start:m_start + SAMPLE] if i not in h and i not in t]
        head += [f"[{s.name}] {i + 1}행: {row_text(s.rows[i])}" for i in h]
        mid += [f"[{s.name}] {i + 1}행: {row_text(s.rows[i])}" for i in m]
        tail += [f"[{s.name}] {i + 1}행: {row_text(s.rows[i])}" for i in t]
        observed = sorted(set(range(min(HEADER_ROWS, n_rows))) | set(h) | set(m) | set(t))
        if n_rows:
            locations.append(f"{s.name}!{ranges([i + 1 for i in observed])}행")
    return ProfileParts(fixed=fixed, locations=locations, head=head, mid=mid, tail=tail)
