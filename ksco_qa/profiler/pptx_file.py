"""PPTX 프로필 (설계 부록 B PPTX 행): 슬라이드 수, 슬라이드별 제목, 표·도표 여부.

관찰 위치는 프로필에 실제로 담은 슬라이드만 적는다. 예산을 넘으면 뒤 슬라이드부터 뺀다.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from ksco_qa.profiler.base import cell_text


@dataclass
class PptxModel:
    slides: list[str]  # "s1: 제목 [표]"
    n_tables: int
    n_charts: int

    def render(self, k: int) -> tuple[str, list[str]]:
        n = len(self.slides)
        locations = [f"슬라이드 1-{k} 제목·표/도표 여부" if k > 1 else "슬라이드 1 제목·표/도표 여부"] if k else []
        lines = [
            f"PPTX · 슬라이드 {n} · 표 있는 슬라이드 {self.n_tables} · 도표 있는 슬라이드 {self.n_charts}",
            "관찰 위치: " + (", ".join(locations) or "없음"),
        ]
        if k < n:
            lines.append(f"(슬라이드 {k + 1}-{n} 미포함)")
        lines.extend(self.slides[:k])
        return "\n".join(lines), locations

    def fit_to(self, budget: int) -> tuple[str, list[str]]:
        k = len(self.slides)
        text, locations = self.render(k)
        while len(text) > budget and k > 0:
            k -= 1
            text, locations = self.render(k)
        return text, locations


def profile(path: Path) -> PptxModel:
    from pptx import Presentation

    prs = Presentation(str(path))
    lines, n_tables, n_charts = [], 0, 0
    for i, slide in enumerate(prs.slides, start=1):
        title = slide.shapes.title.text if slide.shapes.title is not None and slide.shapes.title.has_text_frame else ""
        has_table = any(getattr(s, "has_table", False) and s.has_table for s in slide.shapes)
        has_chart = any(getattr(s, "has_chart", False) and s.has_chart for s in slide.shapes)
        n_tables += has_table
        n_charts += has_chart
        flags = "".join(f" [{name}]" for name, on in (("표", has_table), ("도표", has_chart)) if on)
        lines.append(f"s{i}: {cell_text(title, 60) or '(제목 없음)'}{flags}")
    return PptxModel(lines, n_tables, n_charts)
