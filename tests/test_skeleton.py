"""골격 모듈 import, CLI 명령 구성 (설계 10.1)."""

from __future__ import annotations

import importlib

import pytest
from typer.testing import CliRunner

from ksco_qa.cli import app

MODULES = [
    "ksco_qa.p0",
    "ksco_qa.pipeline",
    "ksco_qa.profiler",
    "ksco_qa.profiler.csv_file",
    "ksco_qa.profiler.spreadsheet",
    "ksco_qa.profiler.pdf",
    "ksco_qa.profiler.document",
    "ksco_qa.profiler.pptx_file",
    "ksco_qa.profiler.hwp",
    "ksco_qa.profiler.zip_file",
    "ksco_qa.profiler.structured",
    "ksco_qa.llm.client",
    "ksco_qa.stages.g0",
    "ksco_qa.stages.g1a",
    "ksco_qa.stages.g1b",
    "ksco_qa.stages.g1c",
    "ksco_qa.stages.g2",
    "ksco_qa.g3.export",
    "ksco_qa.g3.validate",
    "ksco_qa.compare",
]


@pytest.mark.parametrize("module", MODULES)
def test_skeleton_imports(module: str) -> None:
    importlib.import_module(module)


runner = CliRunner()


def test_cli_lists_all_modes() -> None:
    result = runner.invoke(app, ["--help"])
    assert result.exit_code == 0
    for cmd in ("estimate", "sample", "g0-sample", "run", "rerun-failed", "export", "compare"):
        assert cmd in result.output


def test_cli_run_options() -> None:
    result = runner.invoke(app, ["run", "--help"])
    assert result.exit_code == 0
    for opt in ("--stage", "--mode", "--resume", "--llm", "--mock-fault"):
        assert opt in result.output


def test_compare_needs_no_api_key_and_reports_missing_results(no_api_key, tmp_path) -> None:
    result = runner.invoke(app, ["compare", str(tmp_path / "task"), "--with", str(tmp_path / "work.csv")])
    assert result.exit_code == 1
    assert "비교할 03_final.jsonl이 없다" in result.output


def test_estimate_runs_without_api_key(sample_task, no_api_key) -> None:
    import json

    result = runner.invoke(app, ["estimate", str(sample_task)])
    assert result.exit_code == 0, result.output
    report = json.loads(result.output)
    assert report["rows"] > 0
    assert report["total_calls_estimate"] >= report["calls"]["G0_initial"]
    assert report["cost_usd_estimate"] is None
    assert not (sample_task / "03_results").exists()


@pytest.mark.parametrize(
    "args",
    [
        ["run", "t"],
        ["run", "t", "--mode", "batch"],
        ["run", "t", "--stage", "G2"],
        ["sample", "30", "t"],
        ["sample", "30", "t", "--mode", "batch"],
        ["g0-sample", "30", "t"],
        ["rerun-failed", "t"],
    ],
)
def test_model_commands_require_key(no_api_key, args: list[str]) -> None:
    result = runner.invoke(app, args)
    assert result.exit_code == 1
    assert "OPENAI_API_KEY가 설정되지 않음" in result.output


def test_key_present_then_config_checked(monkeypatch: pytest.MonkeyPatch, sample_task) -> None:
    monkeypatch.setenv("OPENAI_API_KEY", "sk-test")
    result = runner.invoke(app, ["run", str(sample_task)])
    assert result.exit_code == 1 and "MODEL_G0" in result.output
    assert not (sample_task / "03_results").exists()


def test_g0_sample_requires_only_g0_model_and_prices() -> None:
    from ksco_qa.config import Config, ConfigError, require_for_model_calls

    cfg = Config(
        MODEL_G0="gpt-6-luna-snapshot",
        PRICE_INPUT=1,
        PRICE_CACHED_INPUT=0.5,
        PRICE_OUTPUT=2,
    )
    require_for_model_calls(cfg, stages=("G0",))
    with pytest.raises(ConfigError, match="MODEL_G1B, MODEL_G2"):
        require_for_model_calls(cfg)


@pytest.mark.parametrize(
    "args",
    [
        ["run", "{t}", "--llm", "mock"],
        ["run", "{t}", "--llm", "mock", "--mode", "batch"],
        ["sample", "3", "{t}", "--llm", "mock"],
        ["rerun-failed", "{t}", "--llm", "mock", "--mock-fault", "schema_fail=0.5"],
    ],
)
def test_mock_needs_no_key_and_reaches_pipeline(no_api_key, sample_task, args: list[str]) -> None:
    result = runner.invoke(app, [a.format(t=sample_task) for a in args])
    assert result.exception is None, result.output
    assert result.exit_code == 0, result.output
    assert "MOCK" in result.output
    run_dir = next((sample_task / "03_results" / "_work").iterdir())
    assert run_dir.name.startswith("mock-")
    # P0→G3까지 실제로 돌고, mock 산출물은 run 폴더 안에만 둔다.
    assert (run_dir / "00_mapping_audit.csv").is_file() and (run_dir / "row_cards.jsonl").is_file()
    assert (run_dir / "g0" / "records.jsonl").is_file()
    assert (run_dir / "g1" / "search.jsonl").is_file()
    assert (run_dir / "g1" / "bundles.jsonl").is_file()
    assert (run_dir / "g1" / "pairs.jsonl").is_file()
    assert (run_dir / "g1" / "rows.jsonl").is_file()
    assert (run_dir / "g2" / "records.jsonl").is_file()
    assert (run_dir / "03_final.jsonl").is_file()
    assert (run_dir / "04_sheet_paste.xlsx").is_file()
    assert (run_dir / "run_report.json").is_file()
    assert sorted(p.name for p in (sample_task / "03_results").iterdir()) == ["_work"]


def test_g0_sample_mock_runs_only_g0_and_writes_review_workbook(no_api_key, sample_task) -> None:
    import json

    from openpyxl import load_workbook

    result = runner.invoke(app, ["g0-sample", "3", str(sample_task), "--llm", "mock"])
    assert result.exit_code == 0, result.output
    run_dir = next((sample_task / "03_results" / "_work").iterdir())
    report = json.loads((run_dir / "g0_sample_report.json").read_text(encoding="utf-8"))
    assert report["sample_n"] == 3
    assert len(report["selected_rows"]) == 3
    assert sum(report["g0_decided_by"].values()) == 3
    assert set(report["g0_decided_by"]) <= {"MODEL", "RULE"}
    assert (run_dir / "g0" / "records.jsonl").is_file()
    calls = [
        json.loads(line)
        for line in (run_dir / "calls.jsonl").read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]
    assert calls and {call["stage"] for call in calls} == {"G0"}
    assert not (run_dir / "g1").exists()
    assert not (run_dir / "g2").exists()

    workbook = load_workbook(run_dir / "g0_review.xlsx", read_only=True, data_only=True)
    sheet = workbook["G0 Review"]
    assert sheet.max_row == 4
    headers = [cell.value for cell in sheet[1]]
    assert headers[-2:] == ["review_verdict", "review_reason"]
    workbook.close()


def test_g0_sample_rejects_request_larger_than_datalist(no_api_key, sample_task) -> None:
    result = runner.invoke(app, ["g0-sample", "99", str(sample_task), "--llm", "mock"])
    assert result.exit_code == 1
    assert "DataList 데이터 행은" in result.output


def test_run_from_g2_reuses_prior_gates(no_api_key, sample_task) -> None:
    initial = runner.invoke(app, ["run", str(sample_task), "--llm", "mock"])
    assert initial.exit_code == 0, initial.output
    [run_dir] = list((sample_task / "03_results" / "_work").iterdir())
    calls_path = run_dir / "calls.jsonl"
    calls_before = calls_path.read_text(encoding="utf-8").splitlines()

    resumed = runner.invoke(
        app,
        ["run", str(sample_task), "--stage", "G2", "--llm", "mock", "--resume", run_dir.name],
    )
    assert resumed.exit_code == 0, resumed.output
    assert calls_path.read_text(encoding="utf-8").splitlines() == calls_before


def test_mock_fault_rejected_without_mock(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("OPENAI_API_KEY", "sk-test")
    result = runner.invoke(app, ["run", "t", "--mock-fault", "schema_fail=0.1"])
    assert result.exit_code == 1 and "--llm mock" in result.output


def test_bad_mock_fault_value(no_api_key, sample_task) -> None:
    result = runner.invoke(app, ["run", str(sample_task), "--llm", "mock", "--mock-fault", "nope=1"])
    assert result.exit_code == 1 and "--mock-fault" in result.output


def test_export_needs_no_key_and_reruns_g3(no_api_key, sample_task) -> None:
    """export: 모델 호출 없이 기존 run의 G3만 다시 (키 불필요)."""
    import json

    from fixtures.gate_records import write_gate_records

    result = runner.invoke(app, ["run", str(sample_task), "--llm", "mock"])
    assert result.exit_code == 0, result.output
    [run_dir] = list((sample_task / "03_results" / "_work").iterdir())
    cards = [json.loads(l) for l in (run_dir / "row_cards.jsonl").read_text(encoding="utf-8").splitlines()]
    write_gate_records(run_dir, cards)
    result = runner.invoke(app, ["export", str(sample_task)])
    assert result.exit_code == 0, result.output
    assert "G3 COMPLETED" in result.output and (run_dir / "03_final.jsonl").is_file()
    assert sorted(p.name for p in (sample_task / "03_results").iterdir()) == ["_work"]


def test_export_without_run_fails(no_api_key, sample_task) -> None:
    result = runner.invoke(app, ["export", str(sample_task)])
    assert result.exit_code == 1 and "run이 없다" in result.output
