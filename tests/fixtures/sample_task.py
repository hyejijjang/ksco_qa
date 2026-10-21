"""작은 가상 Task 폴더를 만든다 (실제 Task 데이터 아님).

저장소에 바이너리를 두지 않으려고 테스트 시점에 tmp 폴더에 생성한다.
담은 경우 (EXPECTED에 행별 기대값):
  MATCHED(CSV·XLSX·PDF·PPTX·HWPX·DOCX·ZIP 멤버 JSON·XML·HTML), 확장자 대소문자 차이, 하위 폴더,
  NOT_FOUND, AMBIGUOUS(디스크+ZIP 같은 이름), CONFLICT(F열 'CSV/PDF'), 여러 형식 전부·일부(PARTIAL_FORMAT),
  HWP(NOT_INSPECTED), D열 끝 공백, G열 빈 행, I열 시트 오류값·"매칭 실패", E열 값 있음(UNRESOLVED),
  수식 셀·하이퍼링크, O열 개인 태깅, R~V 이전 검수 결과 열(모델 입력 금지), 빈 체크박스 행(교훈 #6).
"""

from __future__ import annotations

import csv
import json
import zipfile
from pathlib import Path

import openpyxl

TASK_NAME = "KSCO-99"
HEADERS = [
    "AI_Sending", "id", "직업명", "산출물", "레퍼런스", "출력포맷", "태스크", "검색",
    "직무정보", "비고", "기타", "원본파일", "링크", "비식별화 필", "개인 태깅",
    "Tag", "Q열", "AI 판정", "최종 판정", "Regenerate Task", "원본 Task(백업)", "검수 메모",
]
JOB_ACC = "사무 종사자:경영 관련 사무원:회계 사무원"
JOB_STAT = "사무 종사자：통계‧데이터 관련 사무원：통계 및 설문 조사 사무원"  # 전각 콜론
JOB_PLAIN = "총무 사무원"
INFO_ACC = "회계 사무원은 거래 증빙을 확인하고 장부를 기록하며 월별·분기별 집계표를 작성한다."
INFO_STAT = "통계 및 설문 조사 사무원은 설문 응답을 정리하고 항목별 빈도와 비율을 집계한다."

# 이전 검수 결과 열(R~V)에 넣는 값. row_card·모델 입력에 나오면 안 된다.
LEAK_MARK = "LEAK_PREV_QA"

# (id, 직업, D, E, F, G, I)
ROWS = [
    ("T-001", JOB_ACC, "2024년 월별 매출 현황", None, "CSV",
     "월별 매출 거래 자료를 받아 월 단위 합계와 전월 대비 증감률을 계산한 집계표를 CSV로 작성한다.", INFO_ACC),
    ("T-002", JOB_ACC, "2023년 월별 매출 현황", None, "CSV",
     "2023년 월별 매출 거래 자료를 받아 월 단위 합계와 전월 대비 증감률을 계산한 집계표를 CSV로 작성한다.", INFO_ACC),
    ("T-003", JOB_STAT, "지역별 설문 응답 집계표", None, "XLSX",
     "설문 응답 원자료를 지역별로 분류해 문항별 응답 빈도와 비율을 집계한 표를 XLSX로 작성한다.", INFO_STAT),
    ("T-004", JOB_ACC, "분기 예산 집행 보고서", None, "PDF",
     "분기 예산 배정액과 집행액을 대조해 항목별 집행률과 미집행 사유를 정리한 보고서를 작성한다.", INFO_ACC),
    ("T-005", JOB_ACC, "거래처 대장", None, "xlsx",
     "거래처별 거래 내역을 정리해 거래처 코드·사업자번호·누적 거래액을 담은 대장을 작성한다.", INFO_ACC),
    ("T-006", JOB_STAT, "설문 조사 업무 매뉴얼", "조사 지침서", "PDF",
     "설문 조사 절차 자료를 바탕으로 표본 관리·응답 검토·집계 단계를 정리한 업무 매뉴얼을 작성한다.", INFO_STAT),
    ("T-007", JOB_PLAIN, "월간 업무 보고 ", None, "PPTX",
     "월간 업무 실적 자료를 정리해 부서별 실적과 다음 달 계획을 담은 보고 슬라이드를 작성한다.", "#REF!"),
    ("T-008", JOB_STAT, "통계 보고서", None, "HWPX, XLSX",
     "통계 원자료를 집계해 본문 보고서와 부속 통계표를 작성한다.", INFO_STAT),
    ("T-009", JOB_ACC, "결산 안내문", None, ".DOCX, PDF",
     "결산 일정과 제출 서류를 정리한 안내문을 작성한다.", INFO_ACC),
    ("T-010", JOB_STAT, "설문 원자료", None, "JSON", None, INFO_STAT),
    ("T-011", JOB_STAT, "기관 목록", None, "XML",
     "기관 목록 자료를 정리해 기관 코드와 지역을 담은 목록을 작성한다.", "직무정보 매칭 실패"),
    ("T-012", JOB_PLAIN, "업무 안내 페이지", None, "HTML",
     "업무 안내 자료를 정리해 절차와 문의처 표를 담은 안내 페이지를 작성한다.", None),
    ("T-013", JOB_ACC, "옛 회계 보고서", None, "HWP",
     "회계 보고 자료를 정리해 보고서를 작성한다.", INFO_ACC),
    ("T-014", JOB_STAT, "중복 이름 자료", None, "CSV",
     "자료를 정리해 표를 작성한다.", INFO_STAT),
    ("T-015", JOB_ACC, "형식 미정 자료", None, "CSV/PDF",
     "자료를 정리해 결과를 작성한다.", INFO_ACC),
]
DATA_ROWS = len(ROWS)
TRAILING_CHECKBOX_ROWS = 15

# id → (match_status, inspected, warnings ⊇)
EXPECTED = {
    "T-001": ("MATCHED", True, set()),
    "T-002": ("MATCHED", True, set()),
    "T-003": ("MATCHED", True, set()),
    "T-004": ("NOT_FOUND", False, set()),
    "T-005": ("MATCHED", True, set()),
    "T-006": ("MATCHED", True, set()),
    "T-007": ("MATCHED", True, {"D_TRAILING_SPACE", "JOBINFO_SHEET_ERROR"}),
    "T-008": ("MATCHED", True, set()),
    "T-009": ("MATCHED", True, {"PARTIAL_FORMAT"}),
    "T-010": ("MATCHED", True, {"ORIGINAL_TASK_EMPTY"}),
    "T-011": ("MATCHED", True, {"JOBINFO_MATCH_FAILED"}),
    "T-012": ("MATCHED", True, {"JOBINFO_EMPTY"}),
    "T-013": ("MATCHED", False, set()),
    "T-014": ("AMBIGUOUS", False, set()),
    "T-015": ("CONFLICT", False, set()),
}


def build_sample_task(dest: Path) -> Path:
    """dest 아래에 Task 폴더를 만들고 그 경로를 돌려준다."""
    task = Path(dest)
    out = task / "01_output_files"
    out.mkdir(parents=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DataList"
    ws.append(HEADERS)
    for i, (rid, job, d, e, f, g, info) in enumerate(ROWS, start=2):
        prev_qa = [f"{LEAK_MARK}_{col}_{rid}" for col in "RSTUV"]
        ws.append([False, rid, job, d, e, f, g, "N", info, None, None, None, None, None, "개인태그", None, None, *prev_qa])
        ws[f"J{i}"] = f"=LEN(G{i})"  # 수식 셀 (캐시값 없음)
    ws["M2"].hyperlink = "https://example.com/source"
    ws["M2"].value = "원본 링크"
    for r in range(DATA_ROWS + 2, DATA_ROWS + 2 + TRAILING_CHECKBOX_ROWS):
        ws[f"A{r}"] = False
    other = wb.create_sheet("메모")
    other["A1"] = "다른 탭은 읽지 않는다"
    wb.save(task / f"{TASK_NAME}_DataList.xlsx")

    write_csv(out / "2024년 월별 매출 현황.csv", 2024)
    (out / "매출").mkdir()
    write_csv(out / "매출" / "2023년 월별 매출 현황.csv", 2023, encoding="cp949")
    write_xlsx(out / "지역별 설문 응답 집계표.xlsx")
    write_xlsx(out / "거래처 대장.XLSX")
    write_pdf(out / "설문 조사 업무 매뉴얼.pdf", (
        "1. Overview of the survey procedure and scope",
        "2. Sample management and respondent contact",
        "3. Response review and tabulation steps",
    ))
    write_pptx(out / "월간 업무 보고.pptx")
    write_hwpx(out / "통계 보고서.hwpx")
    write_xlsx(out / "통계 보고서.xlsx")
    write_docx(out / "결산 안내문.docx")  # 결산 안내문.pdf는 없음 → PARTIAL_FORMAT
    (out / "옛 회계 보고서.hwp").write_bytes(b"HWP Document File V3.00 \x1a\x01\x02")
    (out / "중복 이름 자료.csv").write_text("a,b\n1,2\n", encoding="utf-8")
    with zipfile.ZipFile(out / "자료모음.zip", "w") as z:
        z.writestr("원자료/설문 원자료.json", json.dumps(
            {"meta": {"기간": "추후"}, "records": [{"응답자": i, "지역": "서울" if i % 2 else "부산", "점수": i % 5}
                                                 for i in range(1, 8)]}, ensure_ascii=False))
        z.writestr("기관 목록.xml", "<기관목록>" + "".join(
            f"<기관 코드='{i:03d}'><이름>기관{i}</이름><지역>서울</지역></기관>" for i in range(1, 6)) + "</기관목록>")
        z.writestr("업무 안내 페이지.html", "<html><head><title>업무 안내</title></head><body><h1>절차</h1>"
                   "<h2>1. 접수</h2><table><tr><th>구분</th><th>문의처</th></tr><tr><td>회계</td><td>내선 1</td></tr></table>"
                   "</body></html>")
        z.writestr("중복 이름 자료.csv", "a,b\n3,4\n")
    # T-004 "분기 예산 집행 보고서.pdf"는 일부러 없음 (NOT_FOUND)
    return task


def write_csv(path: Path, year: int, encoding: str = "utf-8") -> None:
    with open(path, "w", encoding=encoding, newline="") as f:
        w = csv.writer(f)
        w.writerow(["월", "구분", "매출액", "전월대비증감률"])
        prev = None
        for m in range(1, 13):
            amount = 1000 + m * 37 + (year - 2023) * 50
            rate = "" if prev is None else f"{(amount - prev) / prev:.3f}"
            w.writerow([f"{year}-{m:02d}", "상반기" if m <= 6 else "하반기", amount, rate])
            prev = amount
        w.writerow(["합계", "", "", ""])


def write_xlsx(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "집계"
    ws.merge_cells("A1:C1")
    ws["A1"] = "응답 집계"
    ws.append(["구분", "응답 수", "비율"])
    ws.append(["서울", 120, 0.4])
    ws.append(["부산", 180, 0.6])
    ws.append(["합계", "=SUM(B3:B4)", "=SUM(C3:C4)"])
    wb.save(path)


def write_pdf(path: Path, lines: tuple[str, ...]) -> None:
    """텍스트가 있는 최소 PDF (한 쪽에 한 줄). 내장 폰트만 쓰므로 ASCII 텍스트만 넣는다."""
    ascii_lines = list(lines)
    objects: list[bytes] = [b"<< /Type /Catalog /Pages 2 0 R >>"]
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(len(ascii_lines)))
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(ascii_lines)} >>".encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for i, title in enumerate(ascii_lines):
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 3 0 R >> >> "
            f"/Contents {5 + 2 * i} 0 R >>".encode()
        )
        stream = f"BT /F1 18 Tf 72 760 Td ({title}) Tj ET".encode()
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for no, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % no + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    path.write_bytes(bytes(out))


def write_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for title in ("부서별 실적", "다음 달 계획"):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = title
    slide = prs.slides[0]
    slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1))
    prs.save(path)


def write_hwpx(path: Path) -> None:
    ns = 'xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph"'
    section = (
        f'<?xml version="1.0" encoding="UTF-8"?><hs:sec xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" {ns}>'
        "<hp:p><hp:run><hp:t>1. 조사 개요</hp:t></hp:run></hp:p>"
        "<hp:p><hp:run><hp:t>본문 문단입니다.</hp:t></hp:run></hp:p>"
        "<hp:p><hp:run><hp:tbl><hp:tr><hp:tc><hp:subList><hp:p><hp:run><hp:t>지역</hp:t></hp:run></hp:p></hp:subList></hp:tc>"
        "<hp:tc><hp:subList><hp:p><hp:run><hp:t>응답 수</hp:t></hp:run></hp:p></hp:subList></hp:tc></hp:tr></hp:tbl></hp:run></hp:p>"
        "<hp:p><hp:run><hp:t>2. 결과</hp:t></hp:run></hp:p>"
        "</hs:sec>"
    )
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("mimetype", "application/hwp+zip")
        z.writestr("Contents/section0.xml", section)


def write_docx(path: Path) -> None:
    import docx

    d = docx.Document()
    d.add_heading("결산 일정", level=1)
    d.add_paragraph("결산은 분기 말에 진행한다.")
    d.add_heading("제출 서류", level=1)
    t = d.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text = "서류"
    t.rows[0].cells[1].text = "기한"
    d.save(path)
