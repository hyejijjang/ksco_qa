"""G3 — 최종 레코드·Gate 파일·06 최소 검증·review_queue·run_report (설계 5.7, 9.2 / 06_Result_Schema.md)."""

from __future__ import annotations

import copy
import csv
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from fixtures.gate_records import EXPECTED_FINAL, write_gate_records
from fixtures.sample_task import DATA_ROWS, TASK_NAME
from ksco_qa.config import Config
from ksco_qa.g3 import validate as v
from ksco_qa.g3.export import (
    DUPLICATE_COLUMNS,
    FINAL_JSONL,
    FINAL_OUTPUTS,
    REVISION_COLUMNS,
    VERIFY_COLUMNS,
    review_queue,
    run_g3,
)
from ksco_qa.g3.records import load_inputs
from ksco_qa.p0 import run_p0
from ksco_qa.run import prepare_run

TOP_KEYS = ["schema_version", "dataset", "source_file", "source_sheet", "source_row", "id", "source", "source_assets",
            "revision", "duplicate", "verify", "final"]


@pytest.fixture
def g3_run(sample_task: Path):
    cfg = Config()
    ctx = prepare_run(sample_task, cfg, "mock")
    run_p0(sample_task, cfg, ctx.run_dir, ctx.result_path("00_mapping_audit.csv"))
    write_gate_records(ctx.run_dir, load_inputs(ctx.run_dir).row_cards)
    report = run_g3(ctx, cfg, {"G0": "pv0", "G1B": "pv1", "G2": "pv2"}, source_file=f"{TASK_NAME}_DataList.xlsx")
    records = [json.loads(l) for l in ctx.result_path(FINAL_JSONL).read_text(encoding="utf-8").splitlines()]
    return ctx, report, records


def test_run_completes_and_validation_passes(g3_run) -> None:
    ctx, report, records = g3_run
    assert report["status"] == "COMPLETED", {k: e for k, e in report["validation"].items() if e}
    assert all(not errors for errors in report["validation"].values())
    assert len(records) == DATA_ROWS


def test_outputs_only_inside_mock_run_dir(g3_run, sample_task: Path) -> None:
    ctx, _, _ = g3_run
    for name in FINAL_OUTPUTS:
        assert (ctx.run_dir / name).is_file(), name
    assert sorted(p.name for p in (sample_task / "03_results").iterdir()) == ["_work"]
    assert (ctx.run_dir / "run_report.json").is_file() and (ctx.run_dir / "review_queue.csv").is_file()


def test_output_files_not_left_locked(g3_run) -> None:
    """검증이 xlsx를 read_only로 연 뒤 닫지 않으면 Windows에서 파일이 잠긴다 (옮기기·지우기 실패)."""
    ctx, _, _ = g3_run
    for name in ("00_revision_regenerate.xlsx", "01_duplicate_audit.xlsx", "02_task_verify.xlsx"):
        path = ctx.run_dir / name
        moved = path.with_suffix(".moved")
        path.rename(moved)
        moved.unlink()


def test_final_status_per_rule_path(g3_run) -> None:
    _, _, records = g3_run
    got = {r["id"]: (r["final"]["status"], r["final"]["next_action"]) for r in records}
    assert got == EXPECTED_FINAL


def test_record_shape_follows_06(g3_run) -> None:
    _, _, records = g3_run
    r = next(r for r in records if r["id"] == "T-001")
    assert list(r)[:12] == TOP_KEYS
    assert r["schema_version"] == "ksco-taskqa/canonical-1.0" and r["source_sheet"] == "DataList"
    assert r["dataset"] == TASK_NAME and r["source_file"] == f"{TASK_NAME}_DataList.xlsx"
    assert r["source"]["cells"]["J"]["formula"] == "=LEN(G2)"  # 수식은 source.cells에
    d = r["source_assets"]["deliverable_files"]
    assert d["source_columns_used"] == ["D", "F"] and d["match_method"] == "D_F_EXACT_FILENAME"
    assert len(r["revision"]["g0_self_check"]) == 10 and len(r["verify"]["checks"]) == 16
    check = r["verify"]["checks"]["job_fit"]
    assert check == {"status": "PASS", "evidence": "file_profile 1행",
                     "criterion": "해당 직업의 직무 범위에서 수행 가능한 업무여야 함", "conclusion": "기준 충족 (G2 판정)"}
    assert r["revision"]["effective_task"] == r["revision"]["original_task"]  # KEEP
    assert r["final"]["execution_notes"].startswith("API 파이프라인(MOCK); run_id=mock-")
    for gate in ("revision", "duplicate", "verify", "final"):
        assert r[gate]["decided_by"] in ("MODEL", "RULE", "HUMAN")


def test_multi_format_locator_and_rule_rows(g3_run) -> None:
    _, _, records = g3_run
    by_id = {r["id"]: r for r in records}
    d = by_id["T-008"]["source_assets"]["deliverable_files"]
    assert d["matched_locator"] == "통계 보고서.hwpx | 통계 보고서.xlsx"
    assert d["matched_locators"] == ["통계 보고서.hwpx", "통계 보고서.xlsx"]
    rule = by_id["T-013"]
    assert rule["revision"]["decided_by"] == "RULE" and "NOT_INSPECTED" in rule["revision"]["reason_codes"]
    assert rule["revision"]["g0_self_check"] == {} and rule["verify"]["status"] == "SKIPPED"
    assert by_id["T-011"]["final"]["reason_codes"][0] == "G0_CALL_FAILED"
    assert "ORIGINAL_TASK_EMPTY" in by_id["T-010"]["final"]["reason_codes"]
    assert by_id["T-002"]["final"]["reason"] == "규칙 1: 확정 중복 (대표 행 2, T-001) → 중복 제외"
    assert by_id["T-001"]["revision"]["conclusion"] == "KEEP: " + by_id["T-001"]["revision"]["reason"][:150]
    for rid, code in (("T-012", "G2_CALL_FAILED"), ("T-009", "G2_VALIDATION_FAILED")):
        ver = by_id[rid]["verify"]
        assert (ver["status"], ver["decided_by"], ver["reason_codes"]) == ("REVIEW", "RULE", [code])
        assert by_id[rid]["final"]["reason_codes"] == [code]


def test_gate_xlsx_columns_and_values(g3_run) -> None:
    ctx, _, records = g3_run
    for name, extra in (("00_revision_regenerate.xlsx", REVISION_COLUMNS),
                        ("01_duplicate_audit.xlsx", DUPLICATE_COLUMNS), ("02_task_verify.xlsx", VERIFY_COLUMNS)):
        wb = load_workbook(ctx.run_dir / name)
        assert wb.sheetnames == ["DataList"]
        rows = list(wb.active.iter_rows(values_only=True))
        assert list(rows[0][-len(extra):]) == extra
        assert rows[0][1] == "id" and len(rows) == DATA_ROWS + 1
    dup = list(load_workbook(ctx.run_dir / "01_duplicate_audit.xlsx").active.iter_rows(values_only=True))
    header = dup[0]
    t002 = next(r for r in dup if r[1] == "T-002")
    assert t002[header.index("duplicate_of_id")] == "T-001" and t002[header.index("audit_completeness")] == "COMPLETE"
    t008 = next(r for r in dup if r[1] == "T-008")
    assert t008[header.index("audit_completeness")] == "INCOMPLETE"


def test_run_report_summary_and_seed(g3_run) -> None:
    ctx, report, _ = g3_run
    s = report["summary"]
    assert s["rows"] == DATA_ROWS and s["final_status"] == {"PASS": 3, "FAIL": 3, "REVIEW": 10}
    assert s["g1_candidate_pairs"] == 1 and s["g1_high_overlap_rows"] == 1
    assert report["execution_notes"]["llm"] == "MOCK" and len(report["review_queue"]["seed"]) == 16
    saved = json.loads((ctx.run_dir / "run_report.json").read_text(encoding="utf-8"))
    assert saved["status"] == "COMPLETED"


def test_review_queue_categories(g3_run) -> None:
    ctx, report, records = g3_run
    with open(ctx.run_dir / "review_queue.csv", encoding="utf-8-sig") as f:
        rows = {r["id"]: r for r in csv.DictReader(f)}
    assert "RESELECT_B" in rows["T-005"]["categories"] and "G2_FAIL" in rows["T-007"]["categories"]
    for rid in ("T-011", "T-009", "T-012"):  # G0 호출 실패, G2 검증 2회 실패, G2 호출 2회 실패
        assert "AUTO_CHECK" in rows[rid]["categories"], rid
    assert "DUPLICATE_GROUP_SAMPLE" in rows["T-002"]["categories"]
    assert all(r["review_verdict"] == "" for r in rows.values())
    assert report["review_queue"]["rows"] == len(rows)
    assert review_queue(records, ctx.run_id) == review_queue(records, ctx.run_id)  # 같은 run_id → 같은 표본


def test_overrides_applied_as_human(sample_task: Path) -> None:
    cfg = Config()
    ctx = prepare_run(sample_task, cfg, "mock")
    run_p0(sample_task, cfg, ctx.run_dir, ctx.result_path("00_mapping_audit.csv"))
    cards = load_inputs(ctx.run_dir).row_cards
    write_gate_records(ctx.run_dir, cards)
    row = next(c["source_row"] for c in cards if c["id"] == "T-008")
    overrides = [
        {"source_row": row, "id": "T-008", "gate": "G1", "field": "status", "old_value": "REVIEW", "new_value": "UNIQUE",
         "reviewer": "검토자", "reason": "후보와 다른 업무", "decided_at": "2026-10-22"},
        {"source_row": row, "id": "T-008", "gate": "G1", "field": "search_completed", "old_value": False,
         "new_value": True, "reviewer": "검토자", "reason": "확인 완료", "decided_at": "2026-10-22"},
        {"source_row": row, "id": "T-008", "gate": "G1", "field": "unresolved_candidate_count", "old_value": 1,
         "new_value": 0, "reviewer": "검토자", "reason": "확인 완료", "decided_at": "2026-10-22"},
    ]
    (ctx.run_dir / "overrides.jsonl").write_text("".join(json.dumps(o, ensure_ascii=False) + "\n" for o in overrides),
                                                 encoding="utf-8")
    run_g3(ctx, cfg, {}, source_file="x.xlsx")
    r = next(json.loads(l) for l in ctx.result_path(FINAL_JSONL).read_text(encoding="utf-8").splitlines()
             if json.loads(l)["id"] == "T-008")
    assert r["duplicate"]["status"] == "UNIQUE" and r["duplicate"]["decided_by"] == "HUMAN"
    assert (r["final"]["status"], r["final"]["next_action"]) == ("REVIEW", "REVIEW_MANUAL")  # 이제 G2 REVIEW만 남음
    assert r["final"]["reason_codes"] == ["G2_REVIEW"]


# --- 검사별 실패 사례 ---


def _records(g3_run) -> list[dict]:
    return copy.deepcopy(g3_run[2])


def test_check_regenerate_task_fails(g3_run) -> None:
    recs = _records(g3_run)
    r = next(r for r in recs if r["revision"]["action"] == "REGENERATE")
    r["revision"]["effective_task"] = "다른 Task"
    assert v.check_regenerate_task(recs)


def test_check_duplicate_consistency_fails(g3_run) -> None:
    recs = _records(g3_run)
    r = next(r for r in recs if r["duplicate"]["status"] == "DUPLICATE")
    r["final"]["status"] = "REVIEW"
    r["duplicate"]["representative_id"] = "T-999"
    errors = v.check_duplicate_consistency(recs)
    assert len(errors) == 2


def test_check_search_completed_and_pass_conditions_fail(g3_run) -> None:
    recs = _records(g3_run)
    r = next(r for r in recs if r["duplicate"]["unresolved_candidate_count"])
    r["duplicate"]["search_completed"] = True
    assert v.check_search_completed(recs)
    p = next(r for r in recs if r["final"]["status"] == "PASS")
    p["verify"]["status"] = "REVIEW"
    assert "G2 PASS 아님" in v.check_pass_conditions(recs)[0]


def test_check_match_method_and_reference_fail(g3_run) -> None:
    recs = _records(g3_run)
    recs[0]["source_assets"]["deliverable_files"]["match_method"] = "SIMILAR_NAME"
    assert v.check_match_method(recs)
    r = next(r for r in recs if r["source_assets"]["reference_files"]["mode"] == "GENERATE_REFERENCE")
    r["final"]["reason_codes"] = ["REFERENCE_HOLD"]
    assert v.check_reference_not_auto_review(recs)


def test_check_enums_rejects_unknown_and_null_overlap(g3_run) -> None:
    recs = _records(g3_run)
    recs[0]["verify"]["status"] = "DESIGNABLE"
    recs[1]["duplicate"]["overlap_status"] = None
    recs[2]["final"]["reason_codes"] = ["MADE_UP"]
    errors = v.check_enums(recs)
    assert any("DESIGNABLE" in e for e in errors) and any("overlap_status" in e for e in errors)
    assert any("MADE_UP" in e for e in errors)


def test_check_all_null_fields(g3_run) -> None:
    recs = _records(g3_run)
    for r in recs:
        r["duplicate"]["overlap_status"] = None
    assert v.check_all_null_fields(recs) == ["duplicate.overlap_status가 전 행 null"]
    assert v.check_all_null_fields(_records(g3_run)) == []


def test_check_template_sentences(g3_run) -> None:
    recs = _records(g3_run)
    cards = [{"source_row": r["source_row"], "job": "", "job_leaf": "", "deliverable_title": ""} for r in recs]
    for i, r in enumerate(recs):
        r["revision"]["decided_by"] = "MODEL"
        r["revision"]["reason"] = f"자료 {i}건을 확인함 ('파일{i}')"  # 숫자·따옴표 값만 다른 같은 문장
    errors = v.check_template_sentences(recs, cards)
    assert errors and errors[0].startswith("G0 reason 같은 문장이 16행에 반복")
    assert v.check_template_sentences(_records(g3_run), cards) == []


def test_check_row_identity_and_counts(g3_run) -> None:
    ctx, _, records = g3_run
    inputs = load_inputs(ctx.run_dir)
    swapped = copy.deepcopy(records)
    swapped[0], swapped[1] = swapped[1], swapped[0]
    assert v.check_row_identity(inputs.row_cards, swapped)
    gate = {"01_duplicate_audit.xlsx": ctx.run_dir / "01_duplicate_audit.xlsx"}
    assert v.check_row_counts(len(records), gate, records[:-1])


def test_check_source_cells_detects_changed_value(g3_run) -> None:
    ctx, _, records = g3_run
    inputs = load_inputs(ctx.run_dir)
    changed = copy.deepcopy(records)
    changed[0]["source"]["cells"]["D"]["value"] = "바뀐 값"
    errors = v.check_source_cells(inputs.source_cells, changed, {"02_task_verify.xlsx": ctx.run_dir / "02_task_verify.xlsx"})
    assert any("source.cells가 원본과 다르다" in e for e in errors)
    assert any("02_task_verify.xlsx" in e and "D열" in e for e in errors)


def test_failed_validation_sets_status(g3_run, monkeypatch: pytest.MonkeyPatch) -> None:
    ctx, _, _ = g3_run
    (ctx.run_dir / "g1" / "rows.jsonl").write_text(
        "".join(json.dumps({**json.loads(l), "overlap_status": None}, ensure_ascii=False) + "\n"
                for l in (ctx.run_dir / "g1" / "rows.jsonl").read_text(encoding="utf-8").splitlines()),
        encoding="utf-8")
    report = run_g3(ctx, Config(), {}, source_file="x.xlsx")
    assert report["status"] == "FAILED_VALIDATION"
    assert report["validation"]["all_null_fields"] == ["duplicate.overlap_status가 전 행 null"]
